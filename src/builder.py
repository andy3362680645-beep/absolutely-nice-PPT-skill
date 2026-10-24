"""M1 单页可编辑 PPT 构建与最小 manifest，同步保存独立图层。"""

from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import uuid

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from .config import project_path, sha256_file
from .outline import check_assets
from .text_fit import fit_textframe

ALIGNMENTS = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def write_json(path: Path, value: dict) -> None:
    """在同一目录以临时文件原子更新一份 JSON 运行记录。"""
    temp = path.with_name(path.name + ".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


def normalized(text: str) -> str:
    """去除布局空白，保留文案的其他字符。"""
    return re.sub(r"\s+", "", text)


def _shape_record(shape, identifier: str, role: str, kind: str, order: int) -> dict:
    return {
        "id": identifier, "shape_id": shape.shape_id, "name": shape.name,
        "type": kind, "role": role, "z_order": order,
        "box_emu": [shape.left, shape.top, shape.width, shape.height],
        "box_pt": [value / 12700 for value in (shape.left, shape.top, shape.width, shape.height)],
    }


def verify_reopened(pptx_path: Path, manifest: dict) -> list[str]:
    """逐页核对真实对象；版本 1 不要求新增的多页字段。"""
    issues = []
    prs = Presentation(pptx_path)
    records = manifest["slides"]
    if len(prs.slides) != len(records) or (manifest["schema_version"] == 1 and len(prs.slides) != 1):
        return ["slide_count"]
    if [prs.slide_width, prs.slide_height] != manifest["canvas_emu"]:
        issues.append("canvas_mismatch")
    for index, (slide, record) in enumerate(zip(prs.slides, records), 1):
        page_issues = _verify_shapes(list(slide.shapes), record["shapes"])
        if manifest["schema_version"] == 2:
            if record["slide_index"] != index or slide.slide_id != record["pptx_slide_id"]:
                page_issues.append("slide_identity")
            if slide.notes_slide.notes_text_frame.text != record["notes_text"]:
                page_issues.append("notes_mismatch")
            expected_notes = "\n".join(part for part in (
                record["notes"], "\n".join(f"{source['publisher']}，{source['date']}：{source['url']}" for source in record["sources"])
            ) if part)
            if expected_notes != record["notes_text"] or record["source_ids"] != [item["id"] for item in record["sources"]]:
                page_issues.append("source_association")
            page_issues = [f"page={index}:{record['slide_id']}:{item}" for item in page_issues]
        issues.extend(page_issues)
    return sorted(set(issues))


def _verify_shapes(actual, records):
    """比较一页的独立图片和可编辑文本，保持旧检查语义。"""
    issues = []
    if len(actual) != len(records):
        return ["shape_count"]
    for shape, record in zip(actual, records):
        key = record["id"]
        if shape.shape_id != record["shape_id"] or shape.name != key:
            issues.append(f"shape_identity:{key}")
        if list((shape.left, shape.top, shape.width, shape.height)) != record["box_emu"]:
            issues.append(f"shape_geometry:{key}")
        if record["type"] == "text":
            if not shape.has_text_frame or normalized(shape.text) != normalized(record["text"]):
                issues.append(f"editable_text:{key}")
            else:
                if any(value != Pt(record["margin_pt"]) for value in (shape.text_frame.margin_left, shape.text_frame.margin_right, shape.text_frame.margin_top, shape.text_frame.margin_bottom)):
                    issues.append(f"text_margin:{key}")
                if record.get("fill_color") and (shape.fill.type != 1 or str(shape.fill.fore_color.rgb) != record["fill_color"].upper()):
                    issues.append(f"text_fill:{key}")
                for paragraph in shape.text_frame.paragraphs:
                    if "align" in record and paragraph.alignment != ALIGNMENTS[record["align"]]:
                        issues.append(f"text_alignment:{key}")
                    for run in paragraph.runs:
                        if run.font.name != record["font_name"] or run.font.size != Pt(record["font_size_pt"]):
                            issues.append(f"text_font:{key}")
        elif shape.shape_type != 13 or sha256_file(project_path(record["image_path"])) != record["image_sha256"]:
            issues.append(f"picture_source:{key}")
        elif hashlib.sha256(shape.image.blob).hexdigest() != record["image_sha256"]:
            issues.append(f"picture_embedded:{key}")
    return sorted(set(issues))


def _populate_slide(prs, source_slide, outline, config):
    """按层级构建一页，并返回与该页同步的对象清单。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    records = []
    pictures = sorted(source_slide["images"], key=lambda item: item["role"] != "background")
    for picture in pictures:
        path = project_path(picture["path"])
        x, y, w, h = picture["box_pt"]
        with Image.open(path) as image:
            image_size = list(image.size)
            alpha = list(image.getchannel("A").getextrema()) if "A" in image.getbands() else None
            if picture["role"] != "background":
                scale = min(w / image.width, h / image.height)
                fitted_w, fitted_h = image.width * scale, image.height * scale
                x, y, w, h = x + (w-fitted_w)/2, y + (h-fitted_h)/2, fitted_w, fitted_h
        shape = slide.shapes.add_picture(str(path), Pt(x), Pt(y), Pt(w), Pt(h))
        shape.name = picture["id"]
        record = _shape_record(shape, picture["id"], picture["role"], "picture", len(records))
        record.update({"image_path": picture["path"], "image_sha256": sha256_file(path), "image_size_px": image_size, "alpha_extrema": alpha})
        if "subject_safe_box_pt" in picture:
            record["subject_safe_box_pt"] = picture["subject_safe_box_pt"]
        records.append(record)
    for text in source_slide["texts"]:
        x, y, w, h = text["box_pt"]
        shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        shape.name = text["id"]
        tf = shape.text_frame
        margin = text.get("margin_pt", config["margin_pt"])
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(margin)
        if "fill_color" in text:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(text["fill_color"])
        font = config["fonts"][text["font_key"]]
        size = fit_textframe(shape, text["text"], str(font["resolved_path"]), font["name"], cap_pt=text["cap_pt"], floor_pt=text["floor_pt"], bold=font["bold"])
        if size is None:
            raise ValueError(f"内容超量，请精简: {text['id']}，最低字号 {text['floor_pt']}pt")
        for paragraph in tf.paragraphs:
            paragraph.alignment = ALIGNMENTS[text.get("align", "left")]
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string(text["color"])
        record = _shape_record(shape, text["id"], text["role"], "text", len(records))
        record.update({
            "text": text["text"], "written_text": shape.text,
            "font_key": text["font_key"], "font_name": font["name"], "font_size_pt": size,
            "floor_pt": text["floor_pt"], "bold": font["bold"], "color": text["color"],
            "margin_pt": margin, "effective_box_pt": [x+margin, y+margin, w-2*margin, h-2*margin],
            "font_file_sha256": font["sha256"], "fit_status": "passed",
            "fill_color": text.get("fill_color"), "align": text.get("align", "left"),
        })
        records.append(record)
    sources = outline["sources"] if outline["schema_version"] == 1 else [source for key in source_slide["source_ids"] for source in outline["sources"] if source["id"] == key]
    notes = "\n".join(f"{source['publisher']}，{source['date']}：{source['url']}" for source in sources)
    if outline["schema_version"] == 2:
        notes = "\n".join(part for part in (source_slide["notes"], notes) if part)
    slide.notes_slide.notes_text_frame.text = notes
    record = {"slide_id": source_slide["slide_id"], "slide_index": len(prs.slides),
              "shapes": records, "content_groups": source_slide.get("content_groups", {})}
    if outline["schema_version"] == 2:
        record.update(pptx_slide_id=slide.slide_id, layout_type=source_slide["layout_type"],
                      notes=source_slide["notes"], notes_text=notes, source_ids=source_slide["source_ids"], sources=sources)
    return record


def build_slide(outline: dict, config: dict, out_root: Path) -> dict:
    """保留 M1 schema 1 单页构建。"""
    if outline.get("schema_version") != 1 or len(outline["slides"]) != 1:
        raise ValueError("M1 build_slide 只支持版本 1 单页")
    return _build(outline, config, out_root)


def build_deck(outline: dict, config: dict, out_root: Path) -> dict:
    """构建 schema 2 多页，逐页重开核验后才更新最新运行。"""
    if outline.get("schema_version") != 2:
        raise ValueError("M2 build_deck 只支持版本 2")
    return _build(outline, config, out_root)


def _build(outline, config, out_root):
    """共用版本化构建，不覆盖历史产物。"""
    issues = check_assets(outline)
    if issues:
        raise ValueError("; ".join(issues))
    out_root.mkdir(parents=True, exist_ok=True)
    run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ-") + uuid.uuid4().hex[:8]
    run_dir = out_root / "runs" / run_id
    run_dir.mkdir(parents=True)
    prs = Presentation()
    prs.slide_width = Pt(config["canvas"]["width_pt"])
    prs.slide_height = Pt(config["canvas"]["height_pt"])
    prs.core_properties.title = outline["title"]
    prs.core_properties.author = "PPT_agent"
    slides = []
    for index, source in enumerate(outline["slides"], 1):
        try:
            slides.append(_populate_slide(prs, source, outline, config))
        except ValueError as exc:
            raise ValueError(f"第 {index} 页 [{source['slide_id']}]: {exc}") from exc
    filename = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", outline["title"])[:80].strip(" .") or outline["deck_id"]
    pptx_path = run_dir / f"{filename}_v1_可编辑.pptx"
    prs.save(pptx_path)
    manifest = {
        "schema_version": outline["schema_version"], "deck_id": outline["deck_id"], "run_id": run_id,
        "input_sha256": outline["_input_sha256"], "canvas": config["canvas"],
        "canvas_emu": [prs.slide_width, prs.slide_height], "pptx_file": pptx_path.name,
        "font_records": {key: {k: v for k, v in font.items() if k != "resolved_path"} for key, font in config["fonts"].items()},
        "slides": slides,
    }
    if outline["schema_version"] == 2:
        manifest.update(asset_source_record=outline["asset_source_record"],
                        asset_source_sha256=sha256_file(project_path(outline["asset_source_record"])))
    reopened_issues = verify_reopened(pptx_path, manifest)
    if reopened_issues:
        raise RuntimeError("PPT 重开核对失败: " + "; ".join(reopened_issues))
    write_json(run_dir / "manifest.json", manifest)
    state = {
        "run_id": run_id, "input_sha256": outline["_input_sha256"], "state": "built",
        "pptx_file": pptx_path.name, "pptx_sha256": sha256_file(pptx_path),
        "manifest_sha256": sha256_file(run_dir / "manifest.json"),
        "v1_confirmed": False, "run_dir": (Path("runs") / run_id).as_posix(),
    }
    write_json(out_root / "latest.json", state)
    result = {"passed": True, "run_id": run_id, "pptx": str(pptx_path), "slides": len(slides),
              "shapes": sum(len(slide["shapes"]) for slide in slides)}
    if outline["schema_version"] == 1:
        result["font_sizes"] = {r["id"]: r["font_size_pt"] for r in slides[0]["shapes"] if r["type"] == "text"}
    else:
        result["pages"] = [{"slide_index": slide["slide_index"], "slide_id": slide["slide_id"],
                            "shapes": len(slide["shapes"]), "font_sizes": {r["id"]: r["font_size_pt"] for r in slide["shapes"] if r["type"] == "text"}} for slide in slides]
    return result
