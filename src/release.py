"""构建和验证独立 skill 快照；白名单复制、原创样例、完整性检查与源文件保护。"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from urllib.parse import unquote
import zipfile

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
JSON_RESOURCES = {'.ai/knowledge/visual-profiles.json', '.ai/templates/presentation-audit.json'}
SAMPLE_FILES = ['editable.pptx', *[f'reference-{i:02}.png' for i in range(1, 5)]]
PRIVATE = re.compile(r'(?:[A-Za-z]:[\\/](?:Users|桌面|xwechat_files)[\\/]|[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|sk-[A-Za-z0-9_-]{24,})')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def _source_files():
    return sorted([ROOT / name for name in ['SKILL.md', 'AGENTS.md', 'requirements.txt']]
                  + [p for p in (ROOT / '.ai').rglob('*.md')
                     if not _personal_path(p.relative_to(ROOT).as_posix())]
                  + [ROOT / name for name in JSON_RESOURCES]
                  + list((ROOT / 'src').glob('*.py')))


def _protected_snapshot():
    paths = set(_source_files()) | set(ROOT.glob('*.md')) | {ROOT / 'fonts/README.md'}
    for directory in (ROOT / 'specs', ROOT / 'output/releases'):
        if directory.is_dir():
            paths.update(p for p in directory.rglob('*') if p.is_file() and '__pycache__' not in p.parts
                         and not any(part.startswith('.ppt-agent-stage-') for part in p.relative_to(directory).parts))
    return {str(path.relative_to(ROOT)): sha(path) for path in sorted(paths) if path.is_file()}


def _git_state():
    binary = shutil.which('git')
    if binary is None or not (ROOT / '.git').exists():
        return {'branch': None, 'head': None, 'status': None}
    result = {}
    for key, args in [('branch', ['branch', '--show-current']), ('head', ['rev-parse', 'HEAD']),
                      ('status', ['status', '--porcelain=v1', '--untracked-files=all'])]:
        proc = subprocess.run([binary, *args], cwd=ROOT, capture_output=True, text=True, encoding='utf-8', timeout=30)
        if proc.returncode:
            raise ValueError('无法读取Git状态，未继续发布')
        result[key] = proc.stdout.strip()
    return result


def _portable_docs(package):
    """只生成发行说明；SKILL与产品/视觉规范保持源文件原样。"""
    docs = {
        '使用说明.md': '''# PPT_agent 使用说明

把整个ppt-agent文件夹导入支持本地技能的AI，或请有文件访问能力的AI先读取[SKILL.md](SKILL.md)。开发源、安装副本与发布快照相互独立；更新时保留任务和私有偏好，再换用新副本，不自动覆盖旧目录。

## 制作顺序

主题与受众 → 检索内容资料和真实图片 → 确认准确逐页大纲 → 真实文案灰阶线稿与总览 → 确认密度和布局 → 看参考图选风格 → 原创素材及代表页正式模型渲染 → 确认后扩展整套 → 核对交付。详细规则见[产品流程](.ai/product-sop.md)。可编辑PPT仅按需制作，默认内容修改源为准确大纲。

AI展示实际参考图片供选择，用户也可提供理想图片；AI自动抽象设计参数、组织提示词并原创构图，不要求用户描述风格或写提示词。文案去空泛口号和免责套话，保留学术论证、来源和必要限定。布局与提示词见[线稿指南](.ai/knowledge/wireframe-layouts.md)、[风格指南](.ai/knowledge/style-prompts.md)。

宿主需要能检索、读写资料、查看完整图片和实际生图；交付展示PPT/PDF另需封装工具。本包不附带这些服务、账号或API权限。缺项如实说明；普通导出可以制作线稿和辅助检查，不能冒充最终模型美化。

## 个人偏好

按[偏好协议](.ai/knowledge/user-preferences.md)从[空白模板](.ai/templates/user-preferences.md)建立自己的.local/ppt-agent/preferences.md，或使用指定私有位置。本次要求优先，明确长期偏好才保存并告知；支持查看、修正、删除、暂停和迁移。不能持久读写的宿主提供可携带文本，不宣称自动记忆。个人档案不随公共分发包传播；升级时自行保留，勿填入公共.ai模板后转发。

## 运行辅助工具

使用Python 3.10或更新版本，在本目录安装requirements.txt中的依赖；实际导出还需要LibreOffice和已准备的[白名单字体](fonts/README.md)。本包没有捆绑运行环境和字体。

```powershell
python -m pip install -r requirements.txt
python -m src.presentation_audit --help
python -m src.m1 --help
```

可以用宿主自己的PPT工具制作，再按[src说明](src/README.md)检查成品。新检查器相对输入路径按当前工作目录解析。内附M1 CLI的相对路径按包根目录解析；M2仍为实验代码，自动修改识别与通用整图封装不作为稳定一键能力。

## 示例与复测

已确认的原创四页[可编辑样例](examples/visual-sample/editable.pptx)及四张完整参考PNG位于examples/visual-sample，仅示范可选原生稿能力，不是新默认线稿流程的实测或前置。最终模型原图不在此示例中。复测用[同一任务](.ai/templates/cross-host-trial.md)。

只读入口或包校验不证明已在其他AI实际制作成功。接收方应记录实际平台、工具、产物和未验证项。第三方依赖及字体许可分别适用，本包没有新增开源许可证或转授素材许可。
''',
        'src/README.md': '''# Python辅助工具

运行条件见[使用说明](../使用说明.md)，内容与画面按[产品流程](../.ai/product-sop.md)。本包是独立快照，不依赖原作者目录。默认大纲、线稿、选图及正式生图由宿主执行，以下原生PPT工具按需使用，不是默认前置，不提供自动记忆服务。

## 独立成品检查

```powershell
python -m src.presentation_audit inventory --pptx examples/visual-sample/editable.pptx --out output/inventory.json
python -m src.presentation_audit check --pptx my-report.pptx --map my-report-audit.json --pdf my-report.pdf --out output/audit-01
```

inventory只读取真实对象，角色初始unknown。AI依据确认大纲填写profile_id、require_thanks、每页kind和对象role。角色说明绑定真实PPT哈希，文件变化必须重读。字段示例见[角色模板](../.ai/templates/presentation-audit.json)，不要把模板直接当成真实证据。

字号按[场景配置](../.ai/knowledge/visual-profiles.json)与实际画布换算。复杂继承、旋转组合、未知对象及缺PDF检查如实列出。退出0表示本次客观检查通过，1表示确定问题或输入错误，2表示仍有未验证项；视觉与人类验收始终分开。输出目录拒绝覆盖。

## 旧辅助路径

src.m1保留单页validate/build/check；需要自己准备真实文案、来源、字体和独立素材。src.text_fit是中文适配入口，制作时显式传入场景角色下限，10pt历史技术值不是投影标准。LO/PDF导出需要实际环境条件，不等于模型美化。

m1 prepare-image含早期茶饮、暖白/茶绿样例提示词，不用作其他主题通用渲染器。其他主题由宿主按当前确认页、准确文字与素材组织实际生图。src.m2保留实验多页入口；自动diff、稳定批量模型渲染、通用整图封装尚未因此完成。

## 再次分发

src.release提供build/verify，保持源与发布快照独立。构建要求经过真实核对和人类确认的可选原生样页记录；包内示例清单不冒充新默认流程实测。不要加入自己的任务资料、个人偏好、字体、密钥或缓存后转发。
''',
        'fonts/README.md': '''# 字体准备

本包不含字体。默认可编辑样例使用思源黑体Regular/Bold，接收方准备对应字体后查看或导出。下载入口：[Adobe官方固定版本](https://github.com/adobe-fonts/source-han-sans/releases/tag/2.005R)，许可见官方发布资料。

运行本地测量时可把文件放为fonts/SourceHanSansSC-Regular.otf和fonts/SourceHanSansSC-Bold.otf。PPT家族名、测量文件与实际导出字体应一致，必须核对实际PDF是否替换。

其他白名单与授权要求见[字体知识](../.ai/knowledge/fonts.md)。免费商用与允许分发分别核对，不把字体本体加入本包。字号统一取[场景配置](../.ai/knowledge/visual-profiles.json)，不套用技术默认下限。
''',
        '.ai/knowledge/architecture.md': '''# 发布包能力与结构

从[SKILL.md](../../SKILL.md)进入[产品流程](../product-sop.md)，运行条件见[使用说明](../../使用说明.md)。核心规范与源版本一致，产品任务不用走仓库开发问卷。

默认路径为准确大纲→真实内容线稿→看图选风格→原创素材及正式模型图，原生可编辑稿按需提供。线稿布局、风格提示词和自然文案规则为公共指导；个人偏好协议由宿主在私有文件执行，本包不提供自动记忆数据库。个人实例.local/ppt-agent/preferences.md不在分发包中，公共空白模板可复制使用。

src.presentation_audit读取任意宿主的成品PPT、附加对象角色及可选辅助PDF，按场景配置检查实际字号、结构与可验证边界，未知项明确列出。它不替代视觉/人类验收，不修改旧manifest。src.release提供白名单分发与验证。

M1单页本地辅助路径仍为输入与素材→text_fit→builder→render/checks。m1 prepare-image是早期茶饮提示词，实际生图由宿主会话完成；其他主题按当前确认内容组织。M2为实验多页代码；自动改动识别、稳定批量生图和通用整图封装不作为本包的一键保证。

本包包含.ai/规范、模板和场景JSON，src/辅助代码、字体准备说明，以及原创可编辑样例和预览。不包含Git历史、私人任务、用户参考照片、字体本体或运行环境。跨AI只有真实试跑并核对产物后才登记已验证。
'''}
    for relative, text in docs.items():
        dest = package / relative
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_text(text, encoding='utf-8')


def _slug(text):
    return ''.join(c for c in text.lower().strip() if c in '-_' or not unicodedata.category(c).startswith(('P', 'S'))).replace(' ', '-')


def _personal_path(relative: str) -> bool:
    """约定的本地个人目录不能进入公共包，包括误放公共区的同名目录。"""
    return '.local' in {part.casefold() for part in Path(relative).parts}


def _allowed(relative: str) -> bool:
    if _personal_path(relative):
        return False
    if relative in {'SKILL.md', 'AGENTS.md', 'requirements.txt', '使用说明.md', 'release-manifest.json', 'fonts/README.md'} | JSON_RESOURCES:
        return True
    path = Path(relative)
    if path.parts[0] == '.ai' and path.suffix == '.md':
        return True
    if len(path.parts) == 2 and path.parts[0] == 'src' and path.suffix in {'.py', '.md'}:
        return True
    return relative in {'examples/visual-sample/' + name for name in SAMPLE_FILES + ['sample-manifest.json']}


def verify_release(folder: Path, zip_path: Path, runtime: bool = False) -> dict:
    """核对分发白名单/内容/相对链接；可选在临时提取副本进行运行入口验证。"""
    folder, zip_path = Path(folder).resolve(), Path(zip_path).resolve()
    manifest = json.loads((folder / 'release-manifest.json').read_text(encoding='utf-8'))
    rows = manifest['files']
    expected = {row['path'] for row in rows}
    if len(expected) != len(rows):
        raise ValueError('发布清单存在重复路径')
    actual = {p.relative_to(folder).as_posix() for p in folder.rglob('*') if p.is_file()}
    if actual != expected | {'release-manifest.json'}:
        raise ValueError('发布文件遗漏或混入清单外文件')
    links = 0
    for p in folder.rglob('*'):
        if p.is_symlink():
            raise ValueError('发布包不能包含符号链接')
        if not p.is_file():
            continue
        rel = p.relative_to(folder).as_posix()
        if not _allowed(rel):
            raise ValueError('不允许的发布文件：' + rel)
        if p.suffix in {'.md', '.py', '.json', '.txt'}:
            text = p.read_text(encoding='utf-8')
            if PRIVATE.search(text):
                raise ValueError('发布文件包含私人路径/身份/凭证模式：' + rel)
        if p.suffix == '.md':
            for ref in re.findall(r'\[[^\]]*\]\(([^)\n]+)\)', p.read_text(encoding='utf-8')):
                ref = ref.strip('<>')
                if re.match(r'^[a-z]+://|^mailto:', ref):
                    continue
                loc, _, frag = ref.partition('#')
                target = (p.parent / unquote(loc)).resolve() if loc else p
                if not target.is_relative_to(folder) or not target.exists():
                    raise ValueError('相对引用失效：' + rel + ' → ' + ref)
                if frag and target.suffix == '.md':
                    headings = [_slug(m.group(1)) for m in re.finditer(r'^#{1,6}\s+(.+)$', target.read_text(encoding='utf-8'), re.M)]
                    if unquote(frag) not in headings:
                        raise ValueError('锚点失效：' + rel + ' → ' + ref)
                links += 1
        if p.suffix == '.png':
            with Image.open(p) as image:
                image.verify()
        if p.suffix == '.pptx':
            deck = Presentation(p)
            if len(deck.slides) != 4 or not any(s.has_chart for page in deck.slides for s in page.shapes):
                raise ValueError('示例页数或原生图表不正确')
            if not all(any(s.has_text_frame and s.text.strip() for s in page.shapes) for page in deck.slides):
                raise ValueError('示例缺少可编辑文字')
            with zipfile.ZipFile(p) as pptz:
                for name in pptz.namelist():
                    if name.lower().endswith(('.odttf', '.ttf', '.otf')):
                        raise ValueError('示例不能暗带字体文件')
                    if name.endswith(('.xml', '.rels')) and PRIVATE.search(pptz.read(name).decode('utf-8')):
                        raise ValueError('示例属性/关系/备注含私人身份或路径')
    for row in rows:
        relative = Path(row['path'])
        if relative.is_absolute() or '..' in relative.parts or not (folder / relative).resolve().is_relative_to(folder):
            raise ValueError('发布清单路径非法')
        p = folder / relative
        if sha(p) != row['sha256'] or p.stat().st_size != row['size']:
            raise ValueError('发布哈希或尺寸不一致：' + row['path'])
    with zipfile.ZipFile(zip_path) as archive:
        names = archive.namelist()
        if archive.testzip() is not None or len(names) != len(set(names)):
            raise ValueError('ZIP校验失败')
        if set(names) != {'ppt-agent/' + name for name in actual}:
            raise ValueError('ZIP与文件夹清单不一致')
        for name in names:
            if any(part in {'..', ''} for part in name.split('/')) or '\\' in name or name.startswith('/'):
                raise ValueError('ZIP路径非法')
            if archive.read(name) != (folder / name[len('ppt-agent/'):]).read_bytes():
                raise ValueError('ZIP与文件夹内容不一致')
    smoke = []
    if runtime:
        with tempfile.TemporaryDirectory(prefix='ppt-agent-verify-') as temp:
            with zipfile.ZipFile(zip_path) as archive:
                archive.extractall(temp)
            extracted = Path(temp) / 'ppt-agent'
            env = {**os.environ, 'PYTHONDONTWRITEBYTECODE': '1', 'PYTHONUTF8': '1'}
            commands = [[sys.executable, '-B', '-m', name, '--help'] for name in ['src.m1', 'src.m2', 'src.presentation_audit', 'src.release']]
            commands.append([sys.executable, '-B', '-m', 'src.presentation_audit', 'inventory', '--pptx', 'examples/visual-sample/editable.pptx', '--out', str(Path(temp) / 'inventory.json')])
            for command in commands:
                proc = subprocess.run(command, cwd=extracted, env=env, capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=60)
                if proc.returncode != 0:
                    raise ValueError('隔离运行失败：' + command[3] + '：' + proc.stderr[-400:])
                smoke.append({'module': command[3], 'exit_code': 0})
    return {'passed': True, 'files': len(actual), 'local_links': links, 'zip_integrity': True,
            'privacy_scan': True, 'sample_editable_structure': True, 'runtime_smoke': smoke,
            'runtime_scope': '仅使用当前环境已有依赖验证隔离入口与对象读取；未重跑接收方安装、PDF导出或其他AI制作。'}


def build_release(version: str, sample_run: Path, out_root: Path) -> dict:
    """构建不覆盖的发行快照；只有当前样页客观/视觉/人类确认均具备才纳入包。"""
    if not re.fullmatch(r'[a-z0-9]+(?:-[a-z0-9]+)*', version):
        raise ValueError('版本只能使用小写英文、数字和连字符')
    sample_run, out_root = Path(sample_run).resolve(), Path(out_root).resolve()
    destination = out_root / ('ppt-agent-' + version)
    if destination.exists():
        raise ValueError('版本目录已存在，禁止覆盖')
    record = json.loads((sample_run / 'sample-record.json').read_text(encoding='utf-8'))
    confirmation = record.get('human_v1_confirmation')
    if record.get('technical_status') != 'passed' or record.get('visual_review') != 'passed' or not isinstance(confirmation, dict):
        raise ValueError('样页尚未完成客观、视觉与人类v1确认')
    if confirmation.get('pptx_sha256') != sha(sample_run / 'editable.pptx') or record.get('editable_sha256') != confirmation['pptx_sha256']:
        raise ValueError('样页确认不属于当前PPT')
    listed = {item['path']: item['sha256'] for item in record['files']}
    for name in [*SAMPLE_FILES, 'audit.json']:
        if listed.get(name) != sha(sample_run / name):
            raise ValueError('样页证据改变：' + name)
    audit = json.loads((sample_run / 'audit.json').read_text(encoding='utf-8'))
    if audit['technical_status'] != 'passed' or audit['pptx_sha256'] != confirmation['pptx_sha256']:
        raise ValueError('样页检查证据失配')
    before, git_before = _protected_snapshot(), _git_state()
    out_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.ppt-agent-stage-', dir=out_root) as temp:
        stage = Path(temp)
        package = stage / 'ppt-agent'
        package.mkdir()
        for source in _source_files():
            if source.is_symlink() or not source.is_file():
                raise ValueError('发布源缺失或为符号链接')
            dest = package / source.relative_to(ROOT)
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, dest)
        _portable_docs(package)
        examples = package / 'examples/visual-sample'
        examples.mkdir(parents=True)
        for name in SAMPLE_FILES:
            shutil.copy2(sample_run / name, examples / name)
        write_json(examples / 'sample-manifest.json', {'schema_version': 1, 'kind': 'original-editable-sample',
                   'technical_check': 'passed', 'visual_review': 'passed', 'human_v1_confirmed': True,
                   'files': [{'path': name, 'sha256': sha(examples / name)} for name in SAMPLE_FILES]})
        files = [{'path': p.relative_to(package).as_posix(), 'size': p.stat().st_size, 'sha256': sha(p)}
                 for p in sorted(package.rglob('*')) if p.is_file()]
        write_json(package / 'release-manifest.json', {'schema_version': 1, 'name': 'ppt-agent', 'release': version,
                   'source_revision': git_before['head'], 'source_has_uncommitted_changes': bool(git_before['status']),
                   'includes_runtime': False, 'includes_fonts': False, 'files': files})
        archive_path = stage / f'ppt-agent-{version}.zip'
        with zipfile.ZipFile(archive_path, 'x', compression=zipfile.ZIP_DEFLATED) as archive:
            for p in sorted(package.rglob('*')):
                if p.is_file():
                    archive.write(p, 'ppt-agent/' + p.relative_to(package).as_posix())
        verified = verify_release(package, archive_path)
        if before != _protected_snapshot() or git_before != _git_state():
            raise ValueError('发布期间源文件、受保护任务、旧快照或Git状态变化，未发布')
        write_json(stage / 'build-report.json', {**verified, 'source_and_old_artifacts_unchanged': True,
                                                'branch_head_status_unchanged': True})
        # 保留原子发布边界，失败不会占用正式版本目录。
        stage.rename(destination)
    return {'folder': str(destination / 'ppt-agent'), 'zip': str(destination / f'ppt-agent-{version}.zip'),
            'passed': True, 'source_and_old_artifacts_unchanged': True}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    build = commands.add_parser('build')
    build.add_argument('--version', required=True)
    build.add_argument('--sample-run', type=Path, required=True)
    build.add_argument('--out', type=Path, default=Path('output/releases'))
    verify = commands.add_parser('verify')
    verify.add_argument('--folder', type=Path, required=True)
    verify.add_argument('--zip', type=Path, required=True)
    verify.add_argument('--runtime', action='store_true')
    args = parser.parse_args()
    try:
        result = build_release(args.version, args.sample_run, args.out) if args.command == 'build' else verify_release(args.folder, args.zip, args.runtime)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile, subprocess.SubprocessError) as exc:
        print(json.dumps({'passed': False, 'message': str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
