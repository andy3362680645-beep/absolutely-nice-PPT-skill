"""读取任意宿主的 PPTX，核验实际字号、对象与辅助 PDF；审美与人类验收另行记录。"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import sys

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn

DEFAULT_PROFILES = Path(__file__).resolve().parents[1] / '.ai/knowledge/visual-profiles.json'
TEXT_ROLES = {'title', 'heading', 'body', 'data_label', 'source', 'page_number', 'thanks', 'contact'}
ART_ROLES = {'background', 'illustration', 'decoration'}
PAGE_KINDS = {'cover', 'content', 'appendix', 'thanks'}


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _normalized(text: str) -> str:
    return re.sub(r'\s+', '', text)


def _color(font):
    try:
        return str(font.color.rgb) if font.color.rgb is not None else None
    except (AttributeError, ValueError, TypeError):
        return None


def _frame_units(tf, box: list[float]) -> list[dict]:
    """读取显式 run、段落及当前文本体局部默认字号，不猜母版与主题值。"""
    if not tf.text.strip():
        return []
    styles = []
    for paragraph in tf.paragraphs:
        if not paragraph.text.strip():
            continue
        ppr = paragraph._p.find(qn('a:pPr'))
        paragraph_default = ppr.find(qn('a:defRPr')) if ppr is not None else None
        lst = tf._txBody.find(qn('a:lstStyle'))
        level = lst.find(qn(f'a:lvl{paragraph.level + 1}pPr')) if lst is not None else None
        local_default = level.find(qn('a:defRPr')) if level is not None else None
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            rpr = run._r.find(qn('a:rPr'))
            value, origin = None, 'unresolved'
            for node, label in [(rpr, 'run'), (paragraph_default, 'paragraph'), (local_default, 'text_body')]:
                if node is not None and node.get('sz') is not None:
                    value, origin = int(node.get('sz')) / 100, label
                    break
            styles.append({'text': run.text, 'size_pt': value, 'size_source': origin,
                           'font_name': run.font.name, 'bold': run.font.bold, 'color': _color(run.font),
                           'baseline': int(rpr.get('baseline', '0')) if rpr is not None else 0})
    norm = tf._txBody.find('.//' + qn('a:normAutofit'))
    x, y, w, h = box
    effective = [x + tf.margin_left / 12700, y + tf.margin_top / 12700,
                 w - (tf.margin_left + tf.margin_right) / 12700,
                 h - (tf.margin_top + tf.margin_bottom) / 12700]
    return [{'text': tf.text, 'box_pt': box, 'effective_box_pt': effective, 'styles': styles,
             'autofit': norm is not None, 'font_scale': norm.get('fontScale') if norm is not None else None}]


def _chart_details(chart, box) -> tuple[list[dict], dict]:
    """记录原生图表数据与显式标签字号；特殊类型及未知继承明确保留未知。"""
    styles, parts = [], []
    supported = chart.chart_type in {
        XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE_STACKED,
        XL_CHART_TYPE.LINE_MARKERS_STACKED, XL_CHART_TYPE.LINE_STACKED_100,
        XL_CHART_TYPE.LINE_MARKERS_STACKED_100, XL_CHART_TYPE.BAR_CLUSTERED,
        XL_CHART_TYPE.BAR_STACKED, XL_CHART_TYPE.BAR_STACKED_100,
        XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.COLUMN_STACKED,
        XL_CHART_TYPE.COLUMN_STACKED_100,
    }
    for axis_name in ('category_axis', 'value_axis'):
        try:
            axis = getattr(chart, axis_name)
            font = axis.tick_labels.font
            styles.append({'text': axis_name, 'size_pt': font.size.pt if font.size is not None else None,
                           'size_source': 'chart_axis', 'bold': font.bold, 'font_name': font.name,
                           'color': _color(font), 'baseline': 0})
            if axis.has_title and axis.axis_title.has_text_frame:
                parts.extend(_frame_units(axis.axis_title.text_frame, box))
        except (ValueError, AttributeError):
            supported = False
    if chart.has_title and chart.chart_title.has_text_frame:
        parts.extend(_frame_units(chart.chart_title.text_frame, box))
    if chart.has_legend:
        font = chart.legend.font
        styles.append({'text': 'legend', 'size_pt': font.size.pt if font.size is not None else None,
                       'size_source': 'chart_legend', 'baseline': 0})
    for plot in chart.plots:
        if plot.has_data_labels:
            font = plot.data_labels.font
            styles.append({'text': 'data_labels', 'size_pt': font.size.pt if font.size is not None else None,
                           'size_source': 'chart_data_labels', 'baseline': 0})
    series = []
    try:
        for item in chart.series:
            series.append({'name': item.name, 'values': list(item.values)})
    except (ValueError, AttributeError, TypeError):
        supported = False
    for part in parts:
        styles.extend(part['styles'])
    return ([{'text': '', 'box_pt': box, 'effective_box_pt': box, 'styles': styles,
              'autofit': False, 'font_scale': None, 'chart_labels': True}],
            {'chart_type': int(chart.chart_type), 'supported': supported, 'series': series})


def _objects(shapes, ancestry=()) -> list[dict]:
    objects = []
    for shape in shapes:
        path = [*ancestry, shape.shape_id]
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            objects.extend(_objects(shape.shapes, tuple(path)))
            continue
        box = [float(v / 12700) for v in (shape.left, shape.top, shape.width, shape.height)]
        obj = {'shape_path': path, 'name': shape.name, 'role': 'unknown', 'box_pt': box,
               'rotation': float(shape.rotation), 'coordinate_uncertain': bool(ancestry),
               'kind': 'graphic', 'text': '', 'units': []}
        if shape.has_table:
            obj['kind'] = 'table'
            table = shape.table
            y = box[1]
            for row_index, row in enumerate(table.rows):
                x = box[0]
                for col_index, cell in enumerate(row.cells):
                    if not cell.is_spanned:
                        width = sum(table.columns[i].width for i in range(col_index, col_index + cell.span_width)) / 12700
                        height = sum(table.rows[i].height for i in range(row_index, row_index + cell.span_height)) / 12700
                        obj['units'].extend(_frame_units(cell.text_frame, [x, y, width, height]))
                    x += table.columns[col_index].width / 12700
                y += row.height / 12700
        elif shape.has_chart:
            obj['kind'] = 'chart'
            obj['units'], obj['chart'] = _chart_details(shape.chart, box)
        elif shape.has_text_frame:
            obj['units'] = _frame_units(shape.text_frame, box)
            if obj['units']:
                obj['kind'] = 'text'
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            obj['kind'] = 'picture'
            obj['image_sha256'] = hashlib.sha256(shape.image.blob).hexdigest()
        elif shape.shape_type not in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.FREEFORM}:
            obj['kind'] = 'unsupported'
        obj['text'] = '\n'.join(unit['text'] for unit in obj['units'])
        objects.append(obj)
    return objects


def inventory_presentation(pptx_path: Path) -> dict:
    """只读返回文件哈希、实际页序与叶对象清单；未知角色由使用者依据大纲填写。"""
    pptx_path = Path(pptx_path).resolve()
    prs = Presentation(pptx_path)
    return {'schema_version': 1, 'pptx_sha256': _sha(pptx_path), 'profile_id': 'classroom-diagram',
            'require_thanks': True, 'canvas_pt': [prs.slide_width / 12700, prs.slide_height / 12700],
            'slides': [{'slide_index': i, 'pptx_slide_id': page.slide_id, 'kind': 'content',
                        'message': '', 'objects': _objects(page.shapes)}
                       for i, page in enumerate(prs.slides, 1)]}


def _issue(target: list, code: str, page=None, obj=None, **detail):
    item = {'code': code, **detail}
    if page is not None:
        item['slide_index'] = page
    if obj is not None:
        item['shape_path'] = obj['shape_path']
        item['name'] = obj['name']
    if item not in target:
        target.append(item)


def _valid_profiles(profiles):
    if profiles.get('schema_version') != 1 or profiles.get('base_canvas_pt') != [960, 540]:
        raise ValueError('场景配置版本或基准画布无效')
    for profile in profiles['profiles'].values():
        if not TEXT_ROLES.issubset(profile['roles']):
            raise ValueError('场景缺少文字角色')
        for role in profile['roles'].values():
            numbers = [role['min_pt'], *role['recommended_pt']]
            if len(numbers) != 3 or any(type(n) not in (int, float) or not math.isfinite(n) or n <= 0 for n in numbers):
                raise ValueError('场景字号必须为有效正数')
            if not numbers[0] <= numbers[1] <= numbers[2]:
                raise ValueError('场景字号范围无效')


def _pdf_audit(pdf_path, inventory, annotations, roles, scale, errors, review):
    pymupdf.TOOLS.set_small_glyph_heights(True)
    with pymupdf.open(pdf_path) as doc:
        if len(doc) != len(inventory['slides']):
            _issue(errors, 'pdf_page_count', expected=len(inventory['slides']), observed=len(doc))
            return
        for record, page in zip(inventory['slides'], doc):
            index = record['slide_index']
            w, h = inventory['canvas_pt']
            if abs(page.rect.width - w) > 1 or abs(page.rect.height - h) > 1:
                _issue(errors, 'pdf_canvas', index)
                continue
            slots = []
            for obj in record['objects']:
                role = annotations.get((index, tuple(obj['shape_path'])), 'unknown')
                if obj['coordinate_uncertain'] or obj['rotation'] or role not in roles:
                    continue
                for unit in obj['units']:
                    x, y, width, height = unit['effective_box_pt']
                    if width <= 0 or height <= 0:
                        _issue(errors, 'invalid_text_region', index, obj)
                        continue
                    slots.append({'obj': obj, 'unit': unit, 'role': role,
                                  'rect': pymupdf.Rect(x, y, x + width, y + height), 'observed': []})
            for block in page.get_text('rawdict')['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        for char in span['chars']:
                            if not char['c'].strip():
                                continue
                            rect = pymupdf.Rect(char['bbox'])
                            center = (rect.tl + rect.br) / 2
                            owners = [slot for slot in slots if center in slot['rect']]
                            if len(owners) != 1:
                                _issue(review, 'pdf_unassigned_character', index, character=char['c'])
                                continue
                            slot = owners[0]
                            slot['observed'].append(char['c'])
                            if not (slot['rect'] + (-1, -1, 1, 1)).contains(rect) or not page.rect.contains(rect):
                                _issue(errors, 'pdf_text_overflow', index, slot['obj'])
                            formula = any(style.get('baseline') for style in slot['unit']['styles'])
                            if not formula and span['size'] + .1 < roles[slot['role']]['min_pt'] * scale:
                                _issue(errors, 'pdf_small_font', index, slot['obj'], observed_pt=round(span['size'], 3),
                                       minimum_pt=roles[slot['role']]['min_pt'] * scale)
            for slot in slots:
                if slot['obj']['kind'] != 'chart' and _normalized(''.join(slot['observed'])) != _normalized(slot['unit']['text']):
                    _issue(errors, 'pdf_text_mismatch', index, slot['obj'], expected=slot['unit']['text'], observed=''.join(slot['observed']))
                if slot['obj']['kind'] == 'chart' and not slot['observed']:
                    _issue(review, 'chart_labels_not_observed', index, slot['obj'])


def audit_presentation(pptx_path: Path, mapping: dict, profiles: dict, pdf_path: Path | None = None) -> dict:
    """只读核验真实文件；返回 technical_status，视觉与人类确认始终独立为 pending。"""
    _valid_profiles(profiles)
    actual = inventory_presentation(pptx_path)
    errors, review, warnings = [], [], []
    if mapping.get('schema_version') != 1:
        raise ValueError('角色说明只支持 schema_version=1')
    profile = profiles['profiles'].get(mapping.get('profile_id'))
    if profile is None:
        raise ValueError('角色说明中的场景不存在')
    roles = profile['roles']
    if actual['pptx_sha256'] != mapping.get('pptx_sha256'):
        _issue(errors, 'pptx_hash_mismatch')
    if type(mapping.get('require_thanks')) is not bool:
        raise ValueError('require_thanks 必须明确为 true 或 false')
    pages = mapping.get('slides')
    if not isinstance(pages, list):
        raise ValueError('角色说明缺少页面列表')
    if len(pages) != len(actual['slides']):
        _issue(errors, 'map_page_count')
    w, h = actual['canvas_pt']
    scale = h / 540
    if abs((w / h) / (16 / 9) - 1) > .005:
        _issue(review, 'unsupported_canvas_ratio')
    annotations = {}
    for index, record in enumerate(actual['slides'], 1):
        mapped = pages[index - 1] if index <= len(pages) else {}
        if not isinstance(mapped, dict):
            raise ValueError('页面角色说明必须为对象')
        if mapped.get('slide_index') != index or mapped.get('pptx_slide_id') != record['pptx_slide_id']:
            _issue(errors, 'slide_identity_mismatch', index)
        if mapped.get('kind') not in PAGE_KINDS:
            _issue(review, 'unknown_page_kind', index)
        record['kind'] = mapped.get('kind')
        record['message'] = mapped.get('message', '')
        marks = {}
        for item in mapped.get('objects', []):
            path = item.get('shape_path')
            if not isinstance(path, list) or not path or any(type(n) is not int or n < 1 for n in path):
                raise ValueError('shape_path 必须为正整数列表')
            key = tuple(path)
            if key in marks:
                _issue(errors, 'duplicate_shape_annotation', index, shape_path=path)
            marks[key] = item.get('role', 'unknown')
        live = {tuple(obj['shape_path']) for obj in record['objects']}
        if set(marks) - live:
            _issue(errors, 'shape_identity_mismatch', index)
        for obj in record['objects']:
            role = marks.get(tuple(obj['shape_path']), 'unknown')
            obj['role'] = role
            annotations[(index, tuple(obj['shape_path']))] = role
            if role not in TEXT_ROLES | ART_ROLES:
                _issue(review, 'unmapped_object', index, obj)
            if obj['kind'] in {'text', 'table', 'chart'} and role not in TEXT_ROLES:
                _issue(review, 'text_role_missing', index, obj, text=obj['text'])
            if obj['coordinate_uncertain'] or obj['rotation']:
                _issue(review, 'complex_geometry', index, obj)
            else:
                x, y, width, height = obj['box_pt']
                if min(x, y) < -.1 or width < 0 or height < 0 or x + width > w + .1 or y + height > h + .1:
                    _issue(errors, 'object_out_of_page', index, obj)
            if obj['kind'] == 'unsupported' or (obj['kind'] == 'chart' and not obj['chart']['supported']):
                _issue(review, 'unsupported_object', index, obj)
            if role in roles:
                for unit in obj['units']:
                    if not unit['styles']:
                        _issue(review, 'unknown_font_size', index, obj)
                    for style in unit['styles']:
                        if style.get('baseline'):
                            _issue(review, 'formula_or_baseline', index, obj)
                        elif style['size_pt'] is None:
                            _issue(review, 'unknown_font_size', index, obj)
                        elif style['size_pt'] + .1 < roles[role]['min_pt'] * scale:
                            _issue(errors, 'small_font', index, obj, role=role, observed_pt=style['size_pt'], minimum_pt=roles[role]['min_pt'] * scale)
                    if unit['autofit'] and pdf_path is None:
                        _issue(review, 'autofit_not_rendered', index, obj)
        if record['kind'] == 'thanks':
            for obj in record['objects']:
                if obj['role'] in {'body', 'heading'} or obj['kind'] in {'table', 'chart'}:
                    _issue(errors, 'thanks_mixed_with_content', index, obj)
    if mapping['require_thanks']:
        thanks = [page['slide_index'] for page in actual['slides'] if page['kind'] == 'thanks']
        if thanks != [len(actual['slides'])]:
            _issue(errors, 'thanks_not_separate_last_slide')
    else:
        warnings.append({'code': 'fragment_check', 'message': '本次明确为片段检查，未验收完整汇报的结束页。'})
    if pdf_path is None:
        _issue(review, 'pdf_not_checked')
    else:
        _pdf_audit(Path(pdf_path), actual, annotations, roles, scale, errors, review)
    return {'schema_version': 1, 'technical_status': 'failed' if errors else 'requires_review' if review else 'passed',
            'errors': errors, 'needs_review': review, 'warnings': warnings, 'profile_id': mapping['profile_id'],
            'pptx_sha256': actual['pptx_sha256'], 'pdf_sha256': _sha(Path(pdf_path)) if pdf_path else None,
            'profiles_sha256': hashlib.sha256(json.dumps(profiles, sort_keys=True).encode()).hexdigest(),
            'canvas_pt': actual['canvas_pt'], 'slides': actual['slides'],
            'visual_review': 'pending', 'human_acceptance': 'pending',
            'checked_at_utc': datetime.now(timezone.utc).isoformat()}


MESSAGES = {
    'small_font': '文字小于本场景的阅读下限，请精简或重排后放大。',
    'pdf_small_font': '实际导出的文字偏小，请检查自动缩放。',
    'object_out_of_page': '对象超出页面，请调整位置或尺寸。',
    'pdf_text_overflow': '导出的文字超出可用区域，请重排。',
    'pdf_text_mismatch': 'PPT 与辅助 PDF 的文字未一致，请重新导出并核对。',
    'thanks_mixed_with_content': '谢谢页仍有正文、图表或表格，请单独安排结束页。',
    'thanks_not_separate_last_slide': '缺少唯一且独立的最后一页致谢/交流页。',
    'unknown_font_size': '有继承字号无法可靠确定，需要查看实际页面。',
    'pdf_not_checked': '未提供辅助 PDF，尚未执行渲染后文字几何核验。',
    'unmapped_object': '对象角色未标完整，请依据大纲补齐。',
    'pptx_hash_mismatch': '角色说明属于另一版 PPT，请重新读取当前文件。',
    'complex_geometry': '旋转或分组坐标需要人工核对。',
    'pdf_unassigned_character': '部分字符归属不明确，需要看图核对。',
}


def summary(report):
    labels = {'passed': '本次客观检查通过', 'failed': '发现需要修改的问题', 'requires_review': '还有未验证项目'}
    lines = ['# PPT 检查摘要', '', labels[report['technical_status']], '',
             '本报告检查文件与对象，画面是否协调、角色是否标对及最终效果仍需逐页验收。', '']
    for group, title in [('errors', '需要修改'), ('needs_review', '需要补查')]:
        lines.extend([f'## {title}', ''])
        if not report[group]:
            lines.extend(['无。', ''])
        for item in report[group][:12]:
            where = f"第 {item['slide_index']} 页" if 'slide_index' in item else '整套'
            name = f"（{item['name']}）" if item.get('name') else ''
            lines.append(f"- {where}{name}：{MESSAGES.get(item['code'], '请核对：' + item['code'])}")
        if len(report[group]) > 12:
            lines.append(f"- 另有{len(report[group]) - 12}项，完整定位见audit.json。")
        lines.append('')
    lines.extend(['## 人工看图重点', '', '- 插画是否解释内容、实际主体是否足够大。',
                  '- 图表与解释是否就近对应，数字与图形关系是否准确。',
                  '- 重点色/底纹是否适量、标题与正文是否有层次。',
                  '- 谢谢页是否真正独立，正文是否误标为来源或联系方式。', '',
                  '视觉审查：待进行。人类验收：待确认。', ''])
    return '\n'.join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    for name in ('inventory', 'check'):
        command = sub.add_parser(name)
        command.add_argument('--pptx', type=Path, required=True)
        command.add_argument('--out', type=Path, required=True)
        if name == 'check':
            command.add_argument('--map', type=Path, required=True)
            command.add_argument('--profiles', type=Path, default=DEFAULT_PROFILES)
            command.add_argument('--pdf', type=Path)
    args = parser.parse_args()
    try:
        out = args.out.resolve()
        if out.exists():
            raise ValueError('输出已存在，请使用新的文件或报告目录')
        if args.command == 'inventory':
            result = inventory_presentation(args.pptx)
            out.parent.mkdir(parents=True, exist_ok=True)
            with out.open('x', encoding='utf-8') as f:
                json.dump(result, f, ensure_ascii=False, indent=2)
            print(json.dumps({'inventory': str(out), 'slides': len(result['slides'])}, ensure_ascii=False))
            return 0
        mapping = json.loads(args.map.read_text(encoding='utf-8-sig'))
        profiles = json.loads(args.profiles.read_text(encoding='utf-8-sig'))
        result = audit_presentation(args.pptx, mapping, profiles, args.pdf)
        out.mkdir(parents=True, exist_ok=False)
        (out / 'audit.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        (out / '审查摘要.md').write_text(summary(result), encoding='utf-8')
        print(json.dumps({'technical_status': result['technical_status'], 'errors': len(result['errors']),
                          'needs_review': len(result['needs_review']), 'report': str(out / '审查摘要.md')}, ensure_ascii=False))
        return {'passed': 0, 'failed': 1, 'requires_review': 2}[result['technical_status']]
    except (OSError, ValueError, KeyError, TypeError, AttributeError) as exc:
        print(json.dumps({'technical_status': 'failed', 'message': str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
