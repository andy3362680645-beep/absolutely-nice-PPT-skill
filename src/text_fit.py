"""中文文本适配工具（三道防线的防线①）。

为什么有这个模块：
    python-pptx / power-pptx 自带的 fit_text() 用空格断词（text.split()），
    中文没有空格，整段被当成一个超长单词，会报错或返回"放不下"。
    本模块用 Pillow 逐字符测量 + 按字断行 + 标点避头尾 + 从大到小试字号，
    得到能让文本完整放进文本框的最大字号。

使用方式：
    1. 生成前（防线①）：best_font_size() 计算字号，设置到 run.font.size。
    2. 打开时（防线②）：设置 text_frame.auto_size = TEXT_TO_FIT_SHAPE 兜底。
    3. 渲染后（防线③）：渲染成 PDF 后用 PyMuPDF 做几何检测，见 checks.py。

单位约定：
    72dpi 下 1pt = 1px；pt = EMU / 12700。
"""

from __future__ import annotations

import math
from typing import List, Optional

from PIL import ImageFont

# 行距系数：纯字形高度会低估实际行高，实测 1.2 与 LibreOffice 渲染吻合
LINE_FACTOR = 1.2

# LO 混排的字距可能高于 Pillow，保留 10% 横向余量，防止再次折行。
WIDTH_SAFETY = 0.9

# 默认字号上下限（投影可读性下限 10pt）
DEFAULT_CAP_PT = 24
DEFAULT_FLOOR_PT = 10

# 不允许出现在行首的标点（句号、逗号、右括号、右引号等）
NO_LINE_START = set("，。！？、；：）】》」』”’.,!?;:)]}）%℃‰")
# 不允许出现在行尾的标点（左括号、左引号等）
NO_LINE_END = set("（【《「『“‘([{")

# 英文/数字词中允许出现的字符
_WEST_WORD_CHARS = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_.")


def text_width_px(font: ImageFont.FreeTypeFont, text: str) -> float:
    """返回一行文本在指定字体下的像素宽度。"""
    l, _, r, _ = font.getbbox(text)
    return r - l


def _next_unit(text: str, i: int) -> tuple[str, int]:
    """从位置 i 取下一个排版单位：连续英文数字作为一个词，其余为单字符。"""
    if text[i] in _WEST_WORD_CHARS:
        j = i
        while j < len(text) and text[j] in _WEST_WORD_CHARS:
            j += 1
        return text[i:j], j
    return text[i], i + 1


def _start_new_line(cur: str, unit: str, lines: List[str], font, max_width: float) -> str:
    """结束当前行 cur，把 unit 放到新行；处理超长单位与标点避头尾。"""
    # 行尾避头尾：把左括号/左引号带到下一行
    while cur and cur[-1] in NO_LINE_END:
        unit = cur[-1] + unit
        cur = cur[:-1]

    # 单位比整行还宽（超长英文词）：逐字符硬断
    while unit and text_width_px(font, unit) > max_width:
        seg = ""
        k = 0
        while k < len(unit) and text_width_px(font, seg + unit[k]) <= max_width:
            seg += unit[k]
            k += 1
        if k == 0:
            raise ValueError("文本框宽度无法容纳单个字符")
        if cur:
            lines.append(cur)
            cur = ""
        lines.append(seg)
        unit = unit[k:]

    if cur:
        lines.append(cur)
    cur = unit

    # 行首避头尾：右括号/句号等尽量拉回上一行末尾
    while cur and cur[0] in NO_LINE_START:
        if lines and text_width_px(font, lines[-1] + cur[0]) <= max_width:
            lines[-1] += cur[0]
            cur = cur[1:]
        else:
            break
    return cur


def cjk_wrap(text: str, font: ImageFont.FreeTypeFont, max_width: float) -> List[str]:
    """按像素宽度断行：中文按字断、英文数字尽量保持整词、标点避头尾。

    - 保留显式换行符。
    - 超长英文词（超过一行宽）退化为逐字符硬断。
    - 左括号不挂行尾，右括号/句号不置行首。
    - 换行处不保留多余空格，行首不留空格。
    """
    if not math.isfinite(max_width) or max_width <= 0:
        raise ValueError("文本框可用宽度必须为正数")
    lines: List[str] = []

    for paragraph in text.split("\n"):
        cur = ""
        i = 0
        n = len(paragraph)

        while i < n:
            # 空格：作为单位间分隔；连续空格视作一个
            if paragraph[i] == " ":
                while i < n and paragraph[i] == " ":
                    i += 1
                if i >= n:
                    break
                unit, i = _next_unit(paragraph, i)
                if cur and text_width_px(font, cur + " " + unit) <= max_width:
                    cur += " " + unit
                else:
                    cur = _start_new_line(cur, unit, lines, font, max_width)
                continue

            unit, i = _next_unit(paragraph, i)
            if text_width_px(font, cur + unit) <= max_width:
                cur += unit
            else:
                cur = _start_new_line(cur, unit, lines, font, max_width)

        lines.append(cur)

    return lines


def best_font_size(
    text: str,
    font_path: str,
    max_w_pt: float,
    max_h_pt: float,
    cap_pt: int = DEFAULT_CAP_PT,
    floor_pt: int = DEFAULT_FLOOR_PT,
    line_factor: float = LINE_FACTOR,
    font_index: int = 0,
) -> Optional[int]:
    """返回能让文本完整放进文本框的最大整数字号（pt）。

    参数：
        text: 待排版文本（可含换行）。
        font_path: 用于测量的字体文件（ttf/otf/ttc）。
        max_w_pt: 文本框可用宽度（pt），= (shape.width - 左右内边距)/12700。
        max_h_pt: 文本框可用高度（pt）。
        cap_pt: 字号上限。
        floor_pt: 字号下限（投影可读性下限）。
        line_factor: 行距系数。
        font_index: ttc 集合字体时的字体序号。

    返回：
        最大可用字号；若到 floor_pt 仍放不下，返回 None。
        调用方必须处理 None：精简文案或拆页，禁止静默截断。
    """
    if floor_pt < 1 or cap_pt < floor_pt or not math.isfinite(line_factor) or line_factor <= 0:
        raise ValueError("字号范围或行距系数无效")
    if not all(math.isfinite(value) and value > 0 for value in (max_w_pt, max_h_pt)):
        return None
    for size in range(cap_pt, floor_pt - 1, -1):
        font = ImageFont.truetype(font_path, size, index=font_index)
        try:
            lines = cjk_wrap(text, font, max_w_pt * WIDTH_SAFETY)
        except ValueError:
            continue
        if not lines:
            continue
        if size * line_factor * len(lines) <= max_h_pt:
            return size
    return None


def fit_textframe(
    shape,
    text: str,
    font_path: str,
    font_name: str,
    cap_pt: int = DEFAULT_CAP_PT,
    floor_pt: int = DEFAULT_FLOOR_PT,
    bold: bool = False,
    font_index: int = 0,
):
    """把文本写入 python-pptx 的文本框，并完成防线①+②。

    参数：
        shape: 带 text_frame 的 shape（如 add_textbox 的返回）。
        text: 要写入的文本。
        font_path: 测量/渲染使用的字体文件路径。
        font_name: 写入 PPT 的字体名（必须与 font_path 是同一套字体）。
        cap_pt / floor_pt: 字号上下限。
        bold: 是否加粗。

    返回：
        size：选中的字号；None 表示内容超量，需要精简或拆页（此时文本仍写入，
        并由 normAutofit 兜底，但调用方必须处理该情况）。
    """
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    tf = shape.text_frame
    tf.word_wrap = True
    # 防线②：打开时由 PowerPoint/WPS/LibreOffice 自动缩放兜底
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = MSO_ANCHOR.TOP

    w_pt = (shape.width - tf.margin_left - tf.margin_right) / 12700
    h_pt = (shape.height - tf.margin_top - tf.margin_bottom) / 12700

    size = best_font_size(
        text, font_path, w_pt, h_pt,
        cap_pt=cap_pt, floor_pt=floor_pt, font_index=font_index,
    )
    chosen = size if size is not None else floor_pt

    tf.clear()
    para = tf.paragraphs[0]
    font = ImageFont.truetype(font_path, chosen, index=font_index)
    lines = cjk_wrap(text, font, w_pt * WIDTH_SAFETY) if size is not None else text.split("\n")
    para.text = "\v".join(lines)
    para.space_before = Pt(0)
    para.space_after = Pt(0)
    para.line_spacing = Pt(chosen * LINE_FACTOR)
    for run in para.runs:
        run.font.name = font_name
        run.font.size = Pt(chosen)
        run.font.bold = bold
        props = run._r.get_or_add_rPr()
        east_asian = props.find(qn("a:ea"))
        if east_asian is None:
            east_asian = OxmlElement("a:ea")
            props.append(east_asian)
        east_asian.set("typeface", font_name)
    return size
