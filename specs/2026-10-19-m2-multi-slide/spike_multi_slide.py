"""M2 多页关键路径验证；夹具与证据只写本任务 output。"""

from __future__ import annotations

import argparse
import copy
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import uuid

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt
import pymupdf

from src.config import load_config, sha256_file
from src.outline import check_assets, load_deck_outline, load_outline
from src.render import _session_fonts
from src.checks import audit_page
from src.text_fit import fit_textframe

M1_INPUT = ROOT / "specs/2026-10-18-m1-single-slide/m1-input.json"


def save_json(path, data):
    """保存本次验证证据。"""
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def probe(run):
    """不依赖未来 M2 实现，先真实验证两页及第二页错文。"""
    config = load_config(load_outline(M1_INPUT))
    Image.new("RGB", (1600, 900), "#FAF8F1").save(run / "background.png")
    icon = Image.new("RGBA", (200, 200))
    ImageDraw.Draw(icon).ellipse((10, 10, 190, 190), fill="#647547")
    icon.save(run / "illustration.png")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    slide_records = []
    for index in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(run / "background.png"), 0, 0, Pt(960), Pt(540))
        slide.shapes.add_picture(str(run / "illustration.png"), Pt(730), Pt(240), Pt(140), Pt(140))
        specs = [("title", f"多页技术验证 · 第 {index+1} 页", [40, 40, 860, 70], "bold", 36),
                 ("body", "第一页：验证独立文字与图片。" if index == 0 else "第二页：近 30 款、超 500 万份。", [40, 180, 620, 130], "regular", 24)]
        if index:
            specs.append(("extra", "此行为第二页独有。", [40, 350, 620, 50], "regular", 20))
        records = [{"id": "illustration", "type": "picture", "role": "illustration", "box_pt": [730, 240, 140, 140]}]
        for key, text, box, font_key, cap in specs:
            shape = slide.shapes.add_textbox(*(Pt(v) for v in box))
            shape.name = key
            tf = shape.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            font = config["fonts"][font_key]
            size = fit_textframe(shape, text, str(font["resolved_path"]), font["name"], cap_pt=cap, floor_pt=18, bold=font["bold"])
            assert size is not None
            for paragraph in tf.paragraphs:
                for item in paragraph.runs:
                    item.font.color.rgb = RGBColor.from_string("161A15")
            records.append({"id": key, "type": "text", "text": text, "effective_box_pt": box,
                            "font_key": font_key, "font_size_pt": size, "floor_pt": 18})
        slide_records.append({"slide_id": f"probe-{index+1}", "shapes": records})
    pptx = run / "probe.pptx"
    prs.save(pptx)
    reopened = Presentation(pptx)
    assert len(reopened.slides) == 2
    assert [len(s.shapes) for s in reopened.slides] == [4, 5]
    with _session_fonts(config), tempfile.TemporaryDirectory(prefix="m2-probe-") as profile:
        result = subprocess.run([config["soffice"], "-env:UserInstallation=" + Path(profile).as_uri(),
                                 "--headless", "--convert-to", "pdf", "--outdir", str(run), str(pptx)],
                                capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=45)
    assert result.returncode == 0, result.stdout + result.stderr
    fonts = {key: {k: v for k, v in value.items() if k != "resolved_path"} for key, value in config["fonts"].items()}
    audits, hashes = [], []
    with pymupdf.open(run / "probe.pdf") as doc:
        assert len(doc) == 2
        for index, page in enumerate(doc):
            assert abs(page.rect.width - 960) < .1 and abs(page.rect.height - 540) < .1
            audit = audit_page(page, {"slides": [slide_records[index]], "font_records": fonts})
            assert not audit["issues"], audit["issues"]
            audits.append({"slide_index": index+1, "slide_id": slide_records[index]["slide_id"], **audit})
            png = run / f"reference-{index+1:02d}.png"
            page.get_pixmap(matrix=pymupdf.Matrix(2560/960, 2560/960), alpha=False).save(png)
            hashes.append(sha256_file(png))
        wrong = copy.deepcopy(slide_records[1])
        wrong["shapes"][1]["text"] += "故意增加漏文"
        negative = audit_page(doc[1], {"slides": [wrong], "font_records": fonts})
        assert any(x["code"] == "text_mismatch" for x in negative["issues"])
        assert not audits[0]["issues"] and hashes[0] != hashes[1]
    evidence = {"passed": True, "case": "probe", "slides": 2, "shape_counts": [4, 5],
                "pages": audits, "reference_sha256": hashes,
                "second_page_negative": {"slide_index": 2, "slide_id": "probe-2", "deck_passed": False, "issues": negative["issues"]}}
    save_json(run / "result.json", evidence)
    return {"passed": True, "case": "probe", "slides": 2, "shape_counts": [4, 5], "positive_issues": 0, "second_page_negative_detected": True}


def fixture(run):
    """借用已有原创图片，只在隔离目录保存两页输入与来源夹具。"""
    data = json.loads(M1_INPUT.read_text(encoding="utf-8"))
    data.update(schema_version=2, deck_id="m2-multi-slide", title="M2 两页内部验证")
    second = copy.deepcopy(data["slides"][0])
    first = data["slides"][0]
    first.update(layout_type="cover", notes="第一页独立讲稿。", source_ids=[])
    first["texts"] = [t for t in first["texts"] if t["role"] != "source"]
    first["texts"][0]["text"] = "多页技术验证：第一页"
    second.update(slide_id="slide-02", layout_type="columns", notes="第二页独立讲稿。", source_ids=[data["sources"][0]["id"]])
    second["texts"][0]["text"] = "多页技术验证：第二页"
    data["slides"].append(second)
    provenance = run / "source-record.json"
    data["asset_source_record"] = provenance.relative_to(ROOT).as_posix()
    entries = []
    for picture in first["images"]:
        path = ROOT / picture["path"]
        with Image.open(path) as image:
            entries.append({"path": picture["path"], "sha256": sha256_file(path), "origin": "reused-M1-for-test",
                            "tool": "session image_gen", "model": "unknown", "prompt": "沿用M1已有原创素材；原记录 assets/m1-single-slide/source-record.json",
                            "size_px": list(image.size), "transparent_background": picture["require_alpha"]})
    save_json(provenance, {"schema_version": 1, "images": entries})
    path = run / "input.json"
    save_json(path, data)
    return path, load_deck_outline(path)


def schema_case(run):
    """验证页面身份、来源、第二页素材及换行哈希；不执行转换。"""
    path, data = fixture(run)
    assert not check_assets(data)
    raw = json.loads(path.read_text(encoding="utf-8"))
    negatives = {}
    mutations = {
        "duplicate_slide": lambda d: d["slides"][1].update(slide_id="slide-01"),
        "duplicate_shape": lambda d: d["slides"][1]["texts"][1].update(id="title"),
        "unknown_source": lambda d: d["slides"][1].update(source_ids=["missing"]),
        "out_of_bounds": lambda d: d["slides"][1]["texts"][0].update(box_pt=[900, 0, 200, 50]),
        "nan_box": lambda d: d["slides"][1]["texts"][0].update(box_pt=[float("nan"), 0, 200, 50]),
        "missing_alpha_requirement": lambda d: d["slides"][1]["images"][1].update(require_alpha=False),
    }
    for name, mutate in mutations.items():
        bad = copy.deepcopy(raw)
        mutate(bad)
        target = run / (name + ".json")
        save_json(target, bad)
        try:
            load_deck_outline(target)
        except ValueError as exc:
            assert "第 2 页" in str(exc), str(exc)
            negatives[name] = str(exc)
        else:
            raise AssertionError(name)
    bad = copy.deepcopy(raw)
    bad["slides"][1]["images"][1]["path"] = (run / "missing.png").relative_to(ROOT).as_posix()
    target = run / "missing-image.json"
    save_json(target, bad)
    issues = check_assets(load_deck_outline(target))
    assert any("missing_asset" in x and "page=2:slide-02" in x for x in issues), issues
    negatives["second_page_missing_image"] = issues
    bad = copy.deepcopy(raw)
    bad["slides"][1]["images"][1].update(path=bad["slides"][1]["images"][0]["path"])
    save_json(target, bad)
    issues = check_assets(load_deck_outline(target))
    assert any("alpha_missing" in x and "page=2:slide-02" in x for x in issues), issues
    negatives["second_page_opaque_image"] = issues
    lf_hash = data["_input_sha256"]
    path.write_bytes(path.read_bytes().replace(b"\r\n", b"\n").replace(b"\n", b"\r\n"))
    assert load_deck_outline(path)["_input_sha256"] == lf_hash
    for version in (1, 2):
        bad = copy.deepcopy(raw)
        bad["schema_version"] = version
        save_json(target, bad)
        try:
            load_outline(target)
        except ValueError:
            negatives[f"m1_rejects_multi_version_{version}"] = True
        else:
            raise AssertionError("M1 不应接受多页")
    save_json(run / "result.json", {"passed": True, "case": "schema", "negative_cases": negatives, "crlf_compatible": True})
    return {"passed": True, "case": "schema", "negative_cases": len(negatives), "crlf_compatible": True}


def built_fixture(run):
    """用正式多页组装器建立本次隔离运行。"""
    from src.builder import build_deck
    path, outline = fixture(run)
    config = load_config(outline)
    root = run / "built"
    result = build_deck(outline, config, root)
    state = json.loads((root / "latest.json").read_text(encoding="utf-8"))
    version = root / state["run_dir"]
    manifest = json.loads((version / "manifest.json").read_text(encoding="utf-8"))
    return path, outline, config, root, version, manifest, result


def build_case(run):
    """重开核对必须能拒绝换页、改字、移图和超量。"""
    from src.builder import build_deck, verify_reopened
    _, outline, config, root, version, manifest, result = built_fixture(run)
    pptx = version / manifest["pptx_file"]
    assert result["slides"] == 2 and not verify_reopened(pptx, manifest)
    negatives = {}
    for name in ("swap_pages", "second_page_text", "second_page_picture", "second_page_notes"):
        prs = Presentation(pptx)
        if name == "swap_pages":
            ids = prs.slides._sldIdLst
            ids.insert(0, ids[1])
        elif name == "second_page_text":
            next(s for s in prs.slides[1].shapes if s.has_text_frame).text = "故意改动第二页"
        elif name == "second_page_picture":
            prs.slides[1].shapes[1].left += Pt(5)
        else:
            prs.slides[1].notes_slide.notes_text_frame.text = "故意改变讲稿"
        changed = run / (name + ".pptx")
        prs.save(changed)
        issues = verify_reopened(changed, manifest)
        assert issues, name
        negatives[name] = issues
    latest_before = sha256_file(root / "latest.json")
    bad = copy.deepcopy(outline)
    bad["slides"][1]["texts"][0]["text"] = "超量内容" * 300
    try:
        build_deck(bad, config, root)
    except ValueError as exc:
        assert "第 2 页" in str(exc), str(exc)
        negatives["overflow_preserves_latest"] = str(exc)
    else:
        raise AssertionError("超量未拒绝")
    assert sha256_file(root / "latest.json") == latest_before
    save_json(run / "result.json", {"passed": True, "case": "build", "pages": result["pages"], "negative_cases": negatives})
    return {"passed": True, "case": "build", "slides": 2, "negative_cases": len(negatives)}


def render_case(run):
    """真实多页导出、旧单页契约与错误路径验证。"""
    from src.builder import build_slide
    from src.render import render_deck_reference, render_reference
    _, _, config, _, version, manifest, _ = built_fixture(run)
    result = render_deck_reference(version / manifest["pptx_file"], config, version, manifest["slides"])
    assert len(result["pages"]) == 2
    hashes = []
    for index, page in enumerate(result["pages"], 1):
        assert page["slide_index"] == index and page["slide_id"] == f"slide-{index:02d}"
        image = version / page["reference_png"]
        assert sha256_file(image) == page["reference_sha256"]
        hashes.append(page["reference_sha256"])
    assert hashes[0] != hashes[1]
    try:
        render_deck_reference(version / manifest["pptx_file"], {**config, "soffice": str(run / "missing.exe")}, version, manifest["slides"])
    except ValueError:
        pass
    else:
        raise AssertionError("无效 LO 路径未拒绝")
    single = build_slide(load_outline(M1_INPUT), config, run / "single")
    single_dir = Path(single["pptx"]).parent
    old = render_reference(Path(single["pptx"]), config, single_dir)
    assert Path(old["reference_png"]).name == "reference.png"
    assert sha256_file(single_dir / old["reference_png"]) == old["reference_sha256"]
    save_json(run / "result.json", {"passed": True, "case": "render", "multi": result, "single": old})
    return {"passed": True, "case": "render", "slides": 2, "distinct_reference_images": True, "m1_contract_preserved": True}


def audit_case(run):
    """整套检测覆盖错文、溢出、碰撞、字体、交换页及缺页。"""
    from src.checks import check_deck_v1
    from src.render import render_deck_reference
    _, _, config, _, version, manifest, _ = built_fixture(run)
    pptx = version / manifest["pptx_file"]
    render = render_deck_reference(pptx, config, version, manifest["slides"])
    pdf = version / render["pdf"]
    report = check_deck_v1(pptx, pdf, manifest, config)
    assert report["passed"], report["issues"]
    negatives = {}
    for name in ("second_text", "second_overflow", "second_collision", "second_font"):
        bad = copy.deepcopy(manifest)
        texts = [item for item in bad["slides"][1]["shapes"] if item["type"] == "text"]
        if name == "second_text":
            texts[0]["text"] += "故意漏文"
        elif name == "second_overflow":
            texts[0]["effective_box_pt"][3] = 8
        elif name == "second_collision":
            texts[1]["effective_box_pt"] = list(texts[0]["effective_box_pt"])
        else:
            bad["font_records"]["fake"] = {**bad["font_records"]["regular"], "path": "fonts/Wrong-Regular.otf"}
            texts[1]["font_key"] = "fake"
        failed = check_deck_v1(pptx, pdf, bad, config)
        assert not failed["passed"] and failed["slides"][0]["passed"], (name, failed)
        assert any(item.get("slide_index") == 2 for item in failed["issues"])
        negatives[name] = failed["issues"]
    for name, selection in (("missing_page", [0]), ("swapped_pdf_pages", [1, 0])):
        changed = run / (name + ".pdf")
        with pymupdf.open(pdf) as doc:
            doc.select(selection)
            doc.save(changed)
        failed = check_deck_v1(pptx, changed, manifest, config)
        assert not failed["passed"] and len(failed["slides"]) == 2
        if name == "missing_page":
            assert failed["slides"][1]["issues"][0]["code"] == "pdf_page_missing"
        negatives[name] = failed["issues"]
    save_json(run / "result.json", {"passed": True, "case": "audit", "positive": report, "negative_cases": negatives})
    return {"passed": True, "case": "audit", "slides": 2, "positive_issues": 0, "negative_cases": len(negatives)}


def cli_call(command, input_path, out=None, expected=0, extra=()):
    """以真实子进程运行入口，核对退出码并保留机器可读结果。"""
    args = [sys.executable, "-m", "src.m2", command, "--input", str(input_path)]
    if out is not None:
        args += ["--out", str(out)]
    result = subprocess.run(args + list(extra), cwd=ROOT, capture_output=True, text=True, encoding="utf-8",
                            env={**os.environ, "PYTHONIOENCODING": "utf-8"}, timeout=60)
    assert result.returncode == expected, (args, result.returncode, result.stdout, result.stderr)
    return json.loads((result.stdout or result.stderr).strip().splitlines()[-1])


def cli_case(run):
    """真实三命令、重新检查取消确认、九类证据防护和局部失败汇总。"""
    path, _ = fixture(run)
    root = run / "cli-output"
    assert cli_call("validate", path)["ready"]
    assert not cli_call("validate", path, extra=("--schema-only",))["ready"]
    built = cli_call("build", path, root)
    first = cli_call("check", path, root)
    assert first["passed"] and first["slides"] == 2
    state_path = root / "latest.json"
    state = json.loads(state_path.read_text(encoding="utf-8"))
    # 人类确认字段只作为隔离测试夹具，不表示真实人类已批准页面。
    state.update(v1_confirmed=True, v1_confirmation={"source": "internal-test-fixture"})
    save_json(state_path, state)
    second = cli_call("check", path, root)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    assert second["references"] != first["references"]
    assert not state["v1_confirmed"] and "v1_confirmation" not in state
    raw = json.loads(path.read_text(encoding="utf-8"))
    provenance = ROOT / raw["asset_source_record"]
    original_provenance = provenance.read_bytes()
    negatives = {}
    for name in ("input", "pptx", "manifest", "provenance", "checks", "reference", "run_directory"):
        target = run / ("bad-" + name)
        shutil.copytree(root, target)
        state = json.loads((target / "latest.json").read_text(encoding="utf-8"))
        version = target / state["run_dir"]
        input_path = path
        if name == "input":
            changed = copy.deepcopy(raw)
            changed["slides"][1]["texts"][0]["text"] += "真实内容变化"
            input_path = run / "changed-input.json"
            save_json(input_path, changed)
        elif name in {"pptx", "manifest", "checks", "reference", "provenance"}:
            artifact = {"pptx": version / state["pptx_file"], "manifest": version / "manifest.json",
                        "checks": version / "checks.json", "reference": version / state["render"]["pages"][1]["reference_png"],
                        "provenance": provenance}[name]
            artifact.write_bytes(artifact.read_bytes() + b"\n")
        else:
            state["run_dir"] = "../outside"
            save_json(target / "latest.json", state)
        try:
            failure = cli_call("check", input_path, target, expected=2)
            failed_state = json.loads((target / "latest.json").read_text(encoding="utf-8"))
            assert not failure["passed"] and failed_state["state"] == "failed" and not failed_state["v1_confirmed"]
            negatives[name] = failure
        finally:
            provenance.write_bytes(original_provenance)
    negatives["m1_output_protected"] = cli_call("build", path, ROOT / "output/m1-single-slide", expected=2)
    missing = copy.deepcopy(raw)
    missing["slides"][1]["images"][1]["path"] = (run / "missing.png").relative_to(ROOT).as_posix()
    missing_path = run / "missing-input.json"
    save_json(missing_path, missing)
    missing_result = cli_call("validate", missing_path, expected=2)
    assert any("page=2:slide-02" in x for x in missing_result["issues"])
    negatives["second_page_missing_image"] = missing_result
    collision = copy.deepcopy(raw)
    collision["slides"][1]["texts"][1]["box_pt"] = list(collision["slides"][1]["texts"][0]["box_pt"])
    collision_path = run / "collision-input.json"
    save_json(collision_path, collision)
    failed_root = run / "collision-output"
    cli_call("build", collision_path, failed_root)
    failed = cli_call("check", collision_path, failed_root, expected=2)
    assert failed["pages"][0]["passed"] and not failed["pages"][1]["passed"] and failed["state"] == "failed"
    negatives["second_page_collision"] = failed
    save_json(run / "result.json", {"passed": True, "case": "cli", "build": built, "first_check": first,
                                    "second_check": second, "negative_cases": negatives})
    return {"passed": True, "case": "cli", "slides": 2, "negative_cases": len(negatives), "recheck_clears_confirmation": True}


def main():
    """只提供已经实现的验证分支；失败以非零退出。"""
    parser = argparse.ArgumentParser(description=__doc__)
    cases = {"probe": probe, "schema": schema_case, "build": build_case, "render": render_case, "audit": audit_case, "cli": cli_case}
    parser.add_argument("--case", choices=list(cases), required=True)
    args = parser.parse_args()
    run = ROOT / "output/m2-multi-slide/spike" / (args.case + "-" + uuid.uuid4().hex[:10])
    run.mkdir(parents=True)
    result = cases[args.case](run)
    print(json.dumps({**result, "out": run.relative_to(ROOT).as_posix()}, ensure_ascii=False))


if __name__ == "__main__":
    main()
