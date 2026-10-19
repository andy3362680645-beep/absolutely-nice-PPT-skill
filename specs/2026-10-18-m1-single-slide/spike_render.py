"""M1 单页 PPT、真实 LO 导出、字体与字符几何最小验证。"""

from __future__ import annotations

import argparse
import copy
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import uuid

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from PIL import Image
import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt
from src.text_fit import fit_textframe


def normalized(text: str) -> str:
    """只忽略排版空白，保留数字、标点与近/超等限定词。"""
    return re.sub(r"\s+", "", text)


def audit(page, records: list[dict]) -> list[str]:
    """将字符关联非重叠文本区，核对完整性、字号和边界。"""
    # 标准字框包含字体行高，可能超出可见字形；使用 fontsize 高度字框。
    pymupdf.TOOLS.set_small_glyph_heights(True)
    issues = []
    chars = []
    for block in page.get_text("rawdict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                chars.extend((char, span["size"]) for char in span["chars"])
    rectangles = [pymupdf.Rect(record["rect"]) for record in records]
    for i, rect in enumerate(rectangles):
        for other in rectangles[i + 1:]:
            if (rect & other).get_area() > 1:
                issues.append("text_collision")
    assigned = [[] for _ in records]
    for char, size in chars:
        if not char["c"].strip():
            continue
        box = pymupdf.Rect(char["bbox"])
        center = (box.tl + box.br) / 2
        owners = [i for i, rect in enumerate(rectangles) if center in rect]
        if len(owners) != 1:
            issues.append("unassigned_or_ambiguous_character")
            continue
        owner = owners[0]
        assigned[owner].append(char["c"])
        if not (rectangles[owner] + (-1, -1, 1, 1)).contains(box):
            issues.append("text_overflow")
        if not page.rect.contains(box):
            issues.append("page_overflow")
        if size < 9.9:
            issues.append("font_below_floor")
    for record, actual in zip(records, assigned):
        if normalized("".join(actual)) != normalized(record["text"]):
            issues.append("text_mismatch")
    return sorted(set(issues))


def main() -> None:
    """写独立运行目录，任何转换或核对失败均以非零退出。"""
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=ROOT / "output/m1-single-slide/spike")
    parser.add_argument("--soffice", type=Path, default=Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "LibreOffice/program/soffice.com")
    args = parser.parse_args()
    run = args.out.resolve() / uuid.uuid4().hex[:12]
    run.mkdir(parents=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # 仅验证图片 shape 的最小测试夹具，不作为正式视觉素材。
    Image.new("RGB", (1, 1), "#FAF8F1").save(run / "fixture.png")
    slide.shapes.add_picture(str(run / "fixture.png"), 0, 0, Pt(960), Pt(540))
    specs = [
        ("title", "喜茶 2025 年度营销回顾", (40, 35, 880, 64), 40, True),
        ("subtitle", "用户共创，让灵感成为日常", (40, 112, 880, 48), 28, False),
        ("heading_1", "喜贴回归", (40, 220, 260, 38), 25, True),
        ("body_1", "2025 年 11 月，喜茶 GO 小程序回归喜贴功能。", (40, 275, 260, 130), 22, False),
        ("heading_2", "百万共创", (350, 220, 260, 38), 25, True),
        ("body_2", "上线一周，用户提交画作超 100 万份。", (350, 275, 260, 130), 22, False),
        ("heading_3", "灵感小卡", (660, 220, 260, 38), 25, True),
        ("body_3", "全年发放近 30 款、超 500 万份。", (660, 275, 260, 130), 22, False),
        ("source", "数据来源：北京日报客户端，2026-01-23。", (40, 480, 880, 30), 12, False),
    ]
    records = []
    for name, text, (x, y, w, h), cap, bold in specs:
        shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        shape.name = name
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        weight = "Bold" if bold else "Regular"
        size = fit_textframe(shape, text, str(ROOT / f"fonts/SourceHanSansSC-{weight}.otf"), "思源黑体", cap_pt=cap, bold=bold)
        assert size is not None
        for paragraph in tf.paragraphs:
            for text_run in paragraph.runs:
                text_run.font.color.rgb = RGBColor.from_string("161A15")
        records.append({"id": name, "text": text, "rect": [x, y, x+w, y+h], "font_size": size})
    pptx = run / "probe.pptx"
    prs.save(pptx)
    reopened = Presentation(pptx)
    assert len(reopened.slides) == 1 and len(reopened.slides[0].shapes) == 10
    for shape, record in zip(list(reopened.slides[0].shapes)[1:], records):
        assert normalized(shape.text) == normalized(record["text"])
    with tempfile.TemporaryDirectory(prefix="m1-lo-") as profile:
        result = subprocess.run(
            [str(args.soffice), "-env:UserInstallation=" + Path(profile).as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(run), str(pptx)],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=45,
        )
    assert result.returncode == 0, result.stderr
    pdf = run / "probe.pdf"
    assert pdf.is_file(), result.stdout
    with pymupdf.open(pdf) as doc:
        assert len(doc) == 1
        page = doc[0]
        assert abs(page.rect.width - 960) < 0.1 and abs(page.rect.height - 540) < 0.1
        fonts = [item[3] for item in page.get_fonts()]
        assert fonts and all("SourceHanSansSC" in name for name in fonts), fonts
        assert any("Bold" in name for name in fonts), fonts
        issues = audit(page, records)
        assert not issues, issues
        negative = {}
        altered = copy.deepcopy(records)
        altered[1]["text"] += "故意漏文"
        negative["missing_text"] = "text_mismatch" in audit(page, altered)
        altered = copy.deepcopy(records)
        altered[0]["rect"][3] = 55
        negative["overflow"] = bool(audit(page, altered))
        altered = copy.deepcopy(records)
        altered[1]["rect"] = altered[0]["rect"]
        negative["collision"] = "text_collision" in audit(page, altered)
        assert all(negative.values()), negative
        page.get_pixmap(matrix=pymupdf.Matrix(2560 / page.rect.width, 2560 / page.rect.width), alpha=False).save(run / "reference.png")
    summary = {"passed": True, "slides": 1, "shapes": 10, "fonts": fonts, "geometry_issues": issues, "negative_cases": negative, "out": run.relative_to(ROOT).as_posix(), "records": records}
    (run / "result.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in summary.items() if key != "records"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
