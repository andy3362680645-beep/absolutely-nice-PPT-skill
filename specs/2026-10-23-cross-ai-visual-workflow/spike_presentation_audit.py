"""成品对象读取探针与独立检查器行为验证；夹具不是正式视觉样页。"""
from __future__ import annotations

import argparse
import copy
import json
from pathlib import Path
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt


def fixture(path: Path, unknown: bool = False) -> Path:
    """创建含普通文字、局部继承、表格、二维图表和独立结束页的技术夹具。"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for name, text, y, size in [('title', '成品读取验证', 30, 38), ('body', '中文字号与独立对象', 100, 26)]:
        shape = slide.shapes.add_textbox(Pt(40), Pt(y), Pt(800), Pt(50))
        shape.name = name
        shape.text = text
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    inherited = slide.shapes.add_textbox(Pt(40), Pt(165), Pt(800), Pt(40))
    inherited.name = 'inherited'
    inherited.text = '段落继承字号'
    if not unknown:
        inherited.text_frame.paragraphs[0].font.size = Pt(24)
    table = slide.shapes.add_table(1, 2, Pt(40), Pt(240), Pt(350), Pt(55)).table
    for cell, text in zip(table.rows[0].cells, ['因素', '指标']):
        cell.text = text
        cell.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    data = CategoryChartData()
    data.categories = ['0', '1', '2']
    data.add_series('示例', [2, 3, 6])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Pt(440), Pt(240), Pt(450), Pt(230), data).chart
    chart.category_axis.tick_labels.font.size = Pt(18)
    chart.value_axis.tick_labels.font.size = Pt(18)
    chart.has_legend = False
    thanks = prs.slides.add_slide(prs.slide_layouts[6])
    shape = thanks.shapes.add_textbox(Pt(100), Pt(200), Pt(760), Pt(90))
    shape.name, shape.text = 'thanks', '谢谢大家'
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(48)
    prs.save(path)
    return path


def verify_auditor(out: Path) -> dict:
    """用确定错误和未知状态验证检查器，防止未知样式被默认为通过。"""
    from src.presentation_audit import inventory_presentation, audit_presentation
    profiles = json.loads((ROOT / '.ai/knowledge/visual-profiles.json').read_text(encoding='utf-8'))
    pptx = fixture(out / 'normal.pptx')
    mapping = inventory_presentation(pptx)
    mapping['profile_id'] = 'classroom-diagram'
    mapping['require_thanks'] = True
    for page in mapping['slides']:
        page['kind'] = 'thanks' if page['slide_index'] == 2 else 'content'
        for obj in page['objects']:
            obj['role'] = {'title': 'title', 'body': 'body', 'inherited': 'body', 'thanks': 'thanks'}.get(obj['name'], 'data_label')
    def check(label, current_map, expected, file=pptx, pdf=None):
        result = audit_presentation(file, current_map, profiles, pdf)
        assert result['technical_status'] == expected, (label, result)
        return {'case': label, 'status': expected}
    rows = [check('no_pdf', mapping, 'requires_review')]
    bad = copy.deepcopy(mapping)
    bad['pptx_sha256'] = '0' * 64
    rows.append(check('wrong_hash', bad, 'failed'))
    bad = copy.deepcopy(mapping)
    bad['slides'][0]['objects'].pop()
    rows.append(check('unmapped_object', bad, 'requires_review'))
    bad = copy.deepcopy(mapping)
    bad['slides'][1]['objects'][0]['role'] = 'body'
    rows.append(check('mixed_thanks', bad, 'failed'))
    for label, mutate in [
        ('small_body', lambda s: setattr(s.text_frame.paragraphs[0].runs[0].font, 'size', Pt(12))),
        ('out_of_page', lambda s: setattr(s, 'left', Pt(-40))),
    ]:
        prs = Presentation(pptx)
        mutate(prs.slides[0].shapes[1])
        changed = out / (label + '.pptx')
        prs.save(changed)
        bad = copy.deepcopy(mapping)
        from src.config import sha256_file
        bad['pptx_sha256'] = sha256_file(changed)
        rows.append(check(label, bad, 'failed', file=changed))
    unknown = fixture(out / 'unknown.pptx', unknown=True)
    bad = copy.deepcopy(mapping)
    from src.config import sha256_file
    bad['pptx_sha256'] = sha256_file(unknown)
    rows.append(check('unknown_inheritance', bad, 'requires_review', file=unknown))
    # 真正的 PDF 正例使用强调探针的当前 PPT/PDF，避免伪造渲染证据。
    emphasis = ROOT / 'output/cross-ai-visual-workflow/emphasis-02'
    if (emphasis / 'result.json').is_file():
        record = json.loads((emphasis / 'result.json').read_text(encoding='utf-8'))
        source = emphasis / 'emphasis.pptx'
        m = inventory_presentation(source)
        m.update(profile_id='classroom-diagram', require_thanks=False)
        m['slides'][0]['kind'] = 'content'
        for obj in m['slides'][0]['objects']:
            obj['role'] = 'body' if obj['kind'] == 'text' else 'decoration'
        rows.append(check('real_pdf_positive', m, 'passed', file=source, pdf=emphasis / record['pdf']))
    assert any(row['status'] == 'passed' for row in rows), '缺少真实 PDF 正例'
    return {'passed': True, 'cases': rows}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument('--probe', action='store_true')
    mode.add_argument('--verify-auditor', action='store_true')
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    out = args.out.resolve()
    out.mkdir(parents=True, exist_ok=False)
    if args.verify_auditor:
        result = verify_auditor(out)
    else:
        source = fixture(out / 'probe.pptx', unknown=True)
        prs = Presentation(source)
        p = prs.slides[0].shapes[2].text_frame.paragraphs[0]
        assert p.runs[0].font.size is None and p.font.size is None
        p.font.size = Pt(24)
        assert p.runs[0].font.size is None and p.font.size.pt == 24
        assert prs.slides[0].shapes[3].table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size.pt == 22
        assert prs.slides[0].shapes[4].chart.value_axis.tick_labels.font.size.pt == 18
        group = prs.slides[0].shapes.add_group_shape()
        child = group.shapes.add_textbox(Pt(20), Pt(20), Pt(100), Pt(40))
        child.text = '组合对象'
        assert group.shapes[0].shape_id == child.shape_id
        result = {'passed': True, 'explicit_body_pt': 26, 'paragraph_inheritance_pt': 24,
                  'unknown_inheritance': None, 'table_text_pt': 22, 'chart_label_pt': 18,
                  'group_requires_coordinate_handling': True, 'slides': len(prs.slides)}
    (out / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
