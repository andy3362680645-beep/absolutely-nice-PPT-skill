"""四类原创验证样页：可编辑文字/图表、独立生图素材、真实辅助导出与绑定清单。"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
from pathlib import Path
import shutil
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from src.text_fit import fit_textframe
from src.render import render_deck_reference
from src.presentation_audit import inventory_presentation, audit_presentation, summary

NAVY, TEAL, INK, GOLD = '20395E', '286F83', '263445', 'F5E6A6'
SOURCE = 'https://app.www.gov.cn/govdata/gov/202112/12/479280/article.html'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def add_text(slide, name, content, box, size, role, bold=False, color=INK, align='left', emphasize=None):
    """中文统一通过 text_fit；单行重点以粗体保守测量后保留原生 runs。"""
    shape = slide.shapes.add_textbox(*(Pt(v) for v in box))
    shape.name = f'{role}:{name}'
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(2)
    measurement_bold = bold or emphasize is not None
    weight = 'Bold' if measurement_bold else 'Regular'
    floor = {'title': 32, 'heading': 26, 'body': 22, 'data_label': 18,
             'source': 12, 'page_number': 12, 'thanks': 36, 'contact': 18}[role]
    fitted = fit_textframe(shape, content, str(ROOT / 'fonts' / f'SourceHanSansSC-{weight}.otf'),
                           '思源黑体', cap_pt=size, floor_pt=floor, bold=measurement_bold)
    if fitted is None:
        raise ValueError(f'内容超量，需重排：{name}')
    p = tf.paragraphs[0]
    p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
    if emphasize:
        wrapped = p.text
        if '\v' in wrapped or emphasize not in wrapped:
            raise ValueError(f'局部强调需要单行或独立短句：{name}')
        before, selected, after = wrapped.partition(emphasize)
        for node in list(p._p):
            if node.tag != qn('a:pPr'):
                p._p.remove(node)
        for text, is_bold in [(before, bold), (selected, True), (after, bold)]:
            if not text:
                continue
            run = p.add_run()
            run.text = text
            run.font.name, run.font.size, run.font.bold = '思源黑体', Pt(fitted), is_bold
            ea = OxmlElement('a:ea')
            ea.set('typeface', '思源黑体')
            run._r.get_or_add_rPr().append(ea)
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(color)
    return shape


def marker(slide, name, box):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(Pt(v) for v in box))
    shape.name = 'decoration:' + name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(GOLD)
    shape.line.fill.background()


def picture(slide, path, box, role='illustration'):
    with Image.open(path) as im:
        w, h = im.size
        if role == 'illustration':
            assert 'A' in im.getbands() and im.getchannel('A').getextrema()[0] < 255
    x, y, width, height = box
    scale = min(width / w, height / h)
    actual_w, actual_h = w * scale, h * scale
    shape = slide.shapes.add_picture(str(path), Pt(x + (width - actual_w) / 2), Pt(y + (height - actual_h) / 2),
                                     width=Pt(actual_w), height=Pt(actual_h))
    shape.name = f'{role}:{path.stem}'
    return shape


def page(prs, assets, index, title=None, source=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('F8FBFD')
    picture(slide, assets / 'background.png', [0, 0, 960, 540], 'background')
    if title:
        add_text(slide, 'title', title, [42, 27, 870, 66], 40, 'title', True, NAVY)
    if source:
        add_text(slide, 'source', source, [42, 514, 780, 23], 12, 'source', color='50627A')
    add_text(slide, 'number', f'{index:02}', [866, 514, 48, 23], 12, 'page_number', color='50627A', align='right')
    return slide


def build(brief, assets, out):
    for file in ['concept.png', 'process.png', 'thanks.png', 'background.png', 'generation-record.json']:
        if not (assets / file).is_file():
            raise ValueError('缺少已确认素材：' + file)
    if not brief.is_file():
        raise ValueError('缺少已确认逐页大纲')
    out.mkdir(parents=True, exist_ok=False)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    prs.core_properties.author = 'PPT_agent'
    prs.core_properties.last_modified_by = 'PPT_agent'
    prs.core_properties.title = '把冷链信息讲清楚：四类可编辑样页'
    first = page(prs, assets, 1, '冷链中的环节衔接', '来源：国务院办公厅，2021')
    add_text(first, 'body', '各环节保持衔接，温度信息随过程记录。', [42, 100, 570, 42], 24, 'body', emphasize='各环节保持衔接')
    marker(first, 'key', [634, 104, 280, 31])
    add_text(first, 'key', '过程连续，记录可追溯', [632, 99, 290, 42], 24, 'body', True, NAVY)
    for label, x in [('仓储', 145), ('运输', 432), ('交接', 723)]:
        add_text(first, label, label, [x, 147, 140, 43], 28, 'heading', True, TEAL, 'center')
    picture(first, assets / 'concept.png', [60, 188, 840, 304])
    first.notes_slide.notes_text_frame.text = '概括冷链全过程衔接、监测与追溯的概念。\n来源：国务院办公厅，2021\n' + SOURCE

    second = page(prs, assets, 2, '温度记录的处理流程', '流程：讲解示例')
    add_text(second, 'key', '每一步都有对应记录', [42, 99, 850, 42], 26, 'body', True, TEAL)
    labels = ['发现变化', '核实原因', '记录处置', '回看复盘']
    notes = ['查看记录', '对照现场', '留下依据', '改善衔接']
    for i, (label, note) in enumerate(zip(labels, notes)):
        add_text(second, f'step-{i}', label, [50 + i * 220, 146, 180, 44], 28, 'heading', True, NAVY, 'center')
        add_text(second, f'note-{i}', note, [50 + i * 220, 472, 180, 36], 24, 'body', align='center')
        if i < 3:
            arrow = second.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Pt(235 + i * 220), Pt(158), Pt(28), Pt(15))
            arrow.name = f'illustration:flow-arrow-{i}'
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor.from_string(TEAL)
            arrow.line.fill.background()
    picture(second, assets / 'process.png', [60, 185, 840, 280])
    second.notes_slide.notes_text_frame.text = '讲解用的流程建议，非引用法规中的标准操作程序。'

    third = page(prs, assets, 3, '一段温度记录', '数据：排版演示记录')
    add_text(third, 'y-unit', '温度（℃）', [42, 109, 200, 34], 20, 'data_label', True, TEAL)
    data = CategoryChartData()
    data.categories = ['0', '1', '2', '3', '4']
    data.add_series('温度（℃）', [2, 3, 6, 4, 3])
    frame = third.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Pt(40), Pt(153), Pt(598), Pt(310), data)
    frame.name = 'data_label:temperature-chart'
    chart = frame.chart
    chart.has_legend = chart.has_title = False
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name = '思源黑体'
        axis.tick_labels.font.size = Pt(18)
        axis.tick_labels.font.color.rgb = RGBColor.from_string(INK)
    chart.value_axis.minimum_scale, chart.value_axis.maximum_scale, chart.value_axis.major_unit = 0, 8, 2
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string('DBE7EE')
    series = chart.series[0]
    series.format.line.color.rgb = RGBColor.from_string(TEAL)
    series.format.line.width = Pt(3)
    series.marker.style = XL_MARKER_STYLE.CIRCLE
    series.marker.size = 8
    series.marker.format.fill.solid()
    series.marker.format.fill.fore_color.rgb = RGBColor.from_string(TEAL)
    series.marker.format.line.color.rgb = RGBColor.from_string(TEAL)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.name, plot.data_labels.font.size = '思源黑体', Pt(18)
    plot.data_labels.font.color.rgb = RGBColor.from_string(NAVY)
    plot.data_labels.position = XL_DATA_LABEL_POSITION.ABOVE
    add_text(third, 'x-unit', '时间（小时）', [455, 465, 180, 34], 18, 'data_label', color=TEAL, align='right')
    marker(third, 'peak', [672, 190, 230, 36])
    add_text(third, 'peak', '第 2 小时，\n温度记录升高', [668, 180, 248, 104], 28, 'body', True, NAVY)
    add_text(third, 'explanation', '对照时间点，\n回看现场情况。', [668, 317, 248, 100], 26, 'body')
    third.notes_slide.notes_text_frame.text = '此图为排版演示记录：时间0/1/2/3/4小时，温度2/3/6/4/3℃。不是实测、企业表现或食品温度限值。'

    fourth = page(prs, assets, 4)
    add_text(fourth, 'thanks', '谢谢大家', [68, 181, 430, 95], 54, 'thanks', True, NAVY)
    add_text(fourth, 'exchange', '欢迎交流', [72, 290, 400, 54], 28, 'contact', color=TEAL)
    picture(fourth, assets / 'thanks.png', [500, 103, 390, 375])

    pptx = out / 'editable.pptx'
    prs.save(pptx)
    config = {'canvas': {'width_pt': 960, 'height_pt': 540, 'reference_width_px': 1600},
              'soffice': str(Path(__import__('os').environ.get('ProgramFiles', 'C:/Program Files')) / 'LibreOffice/program/soffice.com'),
              'render_timeout_s': 45,
              'fonts': {key: {'resolved_path': ROOT / 'fonts' / f'SourceHanSansSC-{weight}.otf'}
                        for key, weight in [('regular', 'Regular'), ('bold', 'Bold')]}}
    rendered = render_deck_reference(pptx, config, out, [{'slide_id': f'slide-{i:02}'} for i in range(1, 5)])
    shutil.copy2(out / rendered['pdf'], out / 'reference.pdf')
    for page_record in rendered['pages']:
        shutil.copy2(out / page_record['reference_png'], out / f"reference-{page_record['slide_index']:02}.png")
    mapping = inventory_presentation(pptx)
    for page_record in mapping['slides']:
        index = page_record['slide_index']
        page_record['kind'] = 'thanks' if index == 4 else 'content'
        page_record['message'] = ['环节衔接与记录', '讲解用处理流程', '看清记录变化', '谢谢与交流'][index - 1]
        for obj in page_record['objects']:
            obj['role'] = obj['name'].partition(':')[0]
    (out / 'presentation-audit.json').write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding='utf-8')
    profiles = json.loads((ROOT / '.ai/knowledge/visual-profiles.json').read_text(encoding='utf-8'))
    report = audit_presentation(pptx, mapping, profiles, out / 'reference.pdf')
    (out / 'audit.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    (out / '审查摘要.md').write_text(summary(report), encoding='utf-8')
    (out / 'exact-text.json').write_text(json.dumps([
        {'slide_index': page_record['slide_index'], 'texts': [obj['text'] for obj in page_record['objects'] if obj['text']],
         'chart_data': [obj['chart'] for obj in page_record['objects'] if obj['kind'] == 'chart']}
        for page_record in mapping['slides']], ensure_ascii=False, indent=2), encoding='utf-8')
    files = ['editable.pptx', 'reference.pdf', *[f'reference-{i:02}.png' for i in range(1, 5)],
             'presentation-audit.json', 'exact-text.json', 'audit.json']
    source_record = json.loads((assets / 'generation-record.json').read_text(encoding='utf-8'))
    for item in source_record['assets']:
        item['sha256'] = sha(assets / item['file'])
        with Image.open(assets / item['file']) as im:
            item['size_px'] = list(im.size)
    record = {'schema_version': 1, 'brief_sha256': sha(brief), 'asset_record': source_record,
              'files': [{'path': name, 'sha256': sha(out / name)} for name in files],
              'editable_sha256': sha(pptx), 'technical_status': report['technical_status'],
              'visual_review': 'pending', 'human_v1_confirmation': None, 'model_review': 'pending',
              'render_record': rendered}
    (out / 'sample-record.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    result = {'slides': 4, 'editable': str(pptx), 'technical_status': report['technical_status'],
              'errors': report['errors'], 'needs_review': report['needs_review']}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if report['technical_status'] != 'passed':
        raise SystemExit(1 if report['errors'] else 2)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--brief', type=Path, required=True)
    parser.add_argument('--assets', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    build(args.brief.resolve(), args.assets.resolve(), args.out.resolve())


if __name__ == '__main__':
    main()
