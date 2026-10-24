"""中文局部强调最小验证；使用已登记字体保守测量、真实导出与字符几何核验。"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from src.text_fit import fit_textframe
from src.render import render_reference


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', required=True, type=Path)
    args = parser.parse_args()
    out = args.out.resolve()
    out.mkdir(parents=True, exist_ok=False)
    font = ROOT / 'fonts/SourceHanSansSC-Bold.otf'
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(55), Pt(110), Pt(510), Pt(42))
    marker.fill.solid()
    marker.fill.fore_color.rgb = RGBColor.from_string('F5E6A6')
    marker.line.fill.background()
    text = '重点文字保持可编辑，中文断行完整。'
    box = slide.shapes.add_textbox(Pt(50), Pt(100), Pt(800), Pt(120))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(4)
    size = fit_textframe(box, text, str(font), '思源黑体', cap_pt=30, floor_pt=22, bold=True)
    assert size is not None
    # 只改变现有断行内的 runs，不再以 text_frame.text 覆盖并丢失样式。
    p = tf.paragraphs[0]
    wrapped = p.text
    assert '\v' not in wrapped, '此夹具应单行，跨行场景由正式样页分别测量'
    for element in list(p._p):
        if element.tag != qn('a:pPr'):
            p._p.remove(element)
    for content, bold, color in [('重点文字保持可编辑', True, '20395E'), ('，中文断行完整。', False, '263445')]:
        run = p.add_run()
        run.text = content
        run.font.name, run.font.size, run.font.bold = '思源黑体', Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
        ea = OxmlElement('a:ea')
        ea.set('typeface', '思源黑体')
        run._r.get_or_add_rPr().append(ea)
    source = out / 'emphasis.pptx'
    prs.save(source)
    reopened = Presentation(source)
    actual = reopened.slides[0].shapes[1].text_frame
    assert actual.text == text
    assert [r.font.bold for r in actual.paragraphs[0].runs] == [True, False]
    config = {'soffice': str(Path(__import__('os').environ.get('ProgramFiles', 'C:/Program Files')) / 'LibreOffice/program/soffice.com'),
              'render_timeout_s': 45, 'canvas': {'width_pt': 960, 'height_pt': 540, 'reference_width_px': 1600},
              'fonts': {key: {'resolved_path': ROOT / 'fonts' / f'SourceHanSansSC-{weight}.otf'}
                        for key, weight in [('regular', 'Regular'), ('bold', 'Bold')]}}
    render = render_reference(source, config, out)
    pymupdf.TOOLS.set_small_glyph_heights(True)
    with pymupdf.open(out / render['pdf']) as doc:
        page = doc[0]
        assert ''.join(page.get_text().split()) == ''.join(text.split()), page.get_text()
        region = pymupdf.Rect(54, 104, 846, 216)
        spans = [span for block in page.get_text('rawdict')['blocks'] for line in block.get('lines', []) for span in line['spans']]
        for span in spans:
            assert span['size'] >= 22 - .1
            for char in span['chars']:
                if char['c'].strip():
                    assert (region + (-1, -1, 1, 1)).contains(pymupdf.Rect(char['bbox'])), char
        observed_fonts = sorted(set(span['font'] for span in spans))
        assert all('SourceHanSansSC' in name for name in observed_fonts), observed_fonts
    result = {'passed': True, 'editable_runs': 2, 'size_pt': size, 'full_text_matches': True,
              'glyphs_inside_textbox': True, 'pdf_fonts': observed_fonts, **render}
    (out / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
